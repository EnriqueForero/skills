# -*- coding: utf-8 -*-
"""Generador del libro de declaración de renta (Formulario 210, 18 hojas numeradas).
Uso: python3 generar_excel.py <ruta_caso.py> [salida.xlsx]
El caso es un módulo Python que define CASO (dict) — vea caso_demo.py como esquema completo.
"""
import sys
import datetime as dt
from decimal import Decimal

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

import importlib.util, os
from motor_renta import Params, _tarifa, fila_casilla, FILAS_AUX

RUTA_CASO = sys.argv[1] if len(sys.argv) > 1 else "caso_demo.py"
SALIDA = sys.argv[2] if len(sys.argv) > 2 else None
_spec = importlib.util.spec_from_file_location("caso_usuario", RUTA_CASO)
_mod = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(_mod)
CFG = _mod.CASO
CLAVE = CFG.get("clave", "caso")
BENS = CFG["BENEFICIOS"]
P = Params()
D = Decimal

# ---------------------------------------------------------------- nombres de hoja
SN = {"INICIO": "01_INICIO", "MET": "02_METODOLOGIA", "PERFIL": "03_PERFIL",
      "ING": "04_INGRESOS", "BEN": "05_BENEFICIOS", "DEP": "06_DEPENDIENTES",
      "PAT": "07_PATRIMONIO", "HIST": "08_HISTORICO", "F210": "09_FORMULARIO_210",
      "MX": "10_MOTOR_EXACTO", "LIM": "11_LIMITE_40", "CTRL": "12_CONTROLES",
      "QA": "13_QA", "SOP": "14_SOPORTES", "EXO": "15_EXOGENA",
      "PAR": "16_PARAMETROS", "DEC": "17_DECISIONES", "GLO": "18_GLOSARIO"}
Q = {k: f"'{v}'" for k, v in SN.items()}
PAR, ING, BEN, DEP, PAT, PER = Q["PAR"], Q["ING"], Q["BEN"], Q["DEP"], Q["PAT"], Q["PERFIL"]
F210, MX, LIM, CTRL, QAs, HIST = Q["F210"], Q["MX"], Q["LIM"], Q["CTRL"], Q["QA"], Q["HIST"]
C5 = f"{PAR}!$C$5"

# ---------------------------------------------------------------- estilos
F = "Arial"
NAVY, SUB, GRN, GRNB = "1F3864", "D9E2F3", "E2EFDA", "70AD47"
BLU, AMB = "1F4E79", "FFE699"
ORA_F, ORA_T, RED_F, RED_T, GRIS, TOT = "FCE4D6", "C55A11", "F8CBCC", "9C0006", "808080", "EAF1FA"
TEAL = "1C6E62"
thin = Side(style="thin", color="BFBFBF"); gs = Side(style="thin", color=GRNB)
B_ALL = Border(left=thin, right=thin, top=thin, bottom=thin)
B_GRN = Border(left=gs, right=gs, top=gs, bottom=gs)

def S_title(c): c.font = Font(name=F, size=14, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY)
def S_sec(c): c.font = Font(name=F, size=10, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=NAVY)
def S_sec2(c): c.font = Font(name=F, size=10, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=TEAL)
def S_sub(c): c.font = Font(name=F, size=9, bold=True, color=NAVY); c.fill = PatternFill("solid", fgColor=SUB); c.alignment = Alignment(vertical="center", wrap_text=True)
def S_txt(c, sz=9, b=False, it=False, col="000000"):
    c.font = Font(name=F, size=sz, bold=b, italic=it, color=col); c.alignment = Alignment(vertical="top", wrap_text=True)
def S_in(c, fmt="#,##0"):
    c.font = Font(name=F, size=9); c.fill = PatternFill("solid", fgColor=GRN); c.border = B_GRN
    c.number_format = fmt; c.alignment = Alignment(horizontal="right", vertical="center")
def S_est(c, fmt="#,##0"):
    S_in(c, fmt); c.fill = PatternFill("solid", fgColor=AMB)
def S_fx(c, fmt="#,##0", tot=False, sz=9):
    c.font = Font(name=F, size=sz, bold=True, color=BLU); c.number_format = fmt
    c.alignment = Alignment(horizontal="right", vertical="center")
    if tot: c.fill = PatternFill("solid", fgColor=TOT)
def S_warn(c): c.font = Font(name=F, size=8, color=ORA_T); c.fill = PatternFill("solid", fgColor=ORA_F); c.alignment = Alignment(vertical="top", wrap_text=True)
def S_alert(c): c.font = Font(name=F, size=9, bold=True, color=RED_T); c.fill = PatternFill("solid", fgColor=RED_F); c.alignment = Alignment(vertical="top", wrap_text=True)
def S_note(c): c.font = Font(name=F, size=8, italic=True, color=GRIS); c.alignment = Alignment(vertical="top", wrap_text=True)
def S_cas(c): c.font = Font(name=F, size=9, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2E5395"); c.alignment = Alignment(horizontal="center", vertical="center"); c.border = B_ALL
def S_kpi_l(c): c.font = Font(name=F, size=8, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2E5395"); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
def S_kpi_v(c, fmt='"$" #,##0'):
    c.font = Font(name=F, size=13, bold=True, color=NAVY); c.fill = PatternFill("solid", fgColor=TOT)
    c.number_format = fmt; c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = B_ALL

def put(ws, coord, val, st=None, **kw):
    cell = ws[coord]; cell.value = val
    if st: st(cell, **kw)
    return cell

def widths(ws, spec):
    for col, w in spec.items(): ws.column_dimensions[col].width = w

def header_row(ws, row, cols, labels):
    for col, lab in zip(cols, labels): put(ws, f"{col}{row}", lab, S_sub)

def sec_row(ws, row, txt, cols, st=None):
    put(ws, f"{cols[0]}{row}", txt, st or S_sec)
    for col in cols[1:]: put(ws, f"{col}{row}", "", st or S_sec)

def titulo(ws, cols, txt, sub=""):
    put(ws, "A1", txt, S_title)
    for cc in cols[1:]: put(ws, f"{cc}1", "", S_title)
    last = cols[-1]
    put(ws, f"{last}2", f'=HYPERLINK("#{Q["INICIO"]}!A1","← Volver a 01_INICIO")', S_note)
    if sub:
        put(ws, "A2", sub, S_note)
        ws.merge_cells(f"A2:{cols[-2]}2"); ws.row_dimensions[2].height = 24

wb = openpyxl.Workbook()

# ============================================================ 16_PARAMETROS (creado primero, se reordena al final)
ws = wb.active; ws.title = SN["PAR"]
widths(ws, {"A": 13, "B": 46, "C": 13, "D": 9, "E": 40, "F": 3, "G": 3, "H": 9, "I": 13, "J": 16})
ws.sheet_view.showGridLines = False
titulo(ws, "ABCDE", "PARÁMETROS TRIBUTARIOS — AÑO GRAVABLE 2025 (edite solo si cambia el año o la norma)",
       "Cada valor con su norma. El resto del libro se recalcula desde aquí.")
header_row(ws, 4, "ABCDE", ["Código", "Parámetro", "Valor", "Unidad", "Fuente / norma"])
params = [
    ("UVT_AG", "UVT año gravable 2025", 49799, "COP", "Resolución DIAN 000193 de 2024"),
    ("UVT_PRES", "UVT año de presentación 2026", 52374, "COP", "Resolución DIAN 000238 de 2025"),
    ("LIM_GLOBAL", "Límite conjunto exentas + deducciones (además del 40%)", 1340, "UVT", "Art. 336 num. 3 E.T."),
    ("LIM_25", "Tope renta exenta laboral del 25%", 790, "UVT", "Art. 206 num. 10 E.T."),
    ("LIM_AFC", "Tope AFC + FVP + AVC (y 30% del ingreso del año)", 3800, "UVT", "Arts. 126-1 y 126-4 E.T."),
    ("LIM_VIV", "Intereses de vivienda (tope anual conjunto)", 1200, "UVT", "Art. 119 E.T."),
    ("LIM_ICETEX", "Intereses ICETEX", 100, "UVT", "Art. 119 inc. 3 E.T."),
    ("LIM_SALUD_MES", "Medicina prepagada (por mes)", 16, "UVT/mes", "Art. 387 E.T."),
    ("LIM_DEP_MES", "Dependientes art. 387 (por mes)", 32, "UVT/mes", "Art. 387 inc. 2 E.T."),
    ("DEP_ADIC", "Dependiente adicional fuera del límite", 72, "UVT/año", "Art. 336 num. 3 inc. 2 E.T."),
    ("MAX_DEP", "Máximo dependientes de 72 UVT", 4, "pers.", "Art. 336 num. 3 inc. 2 E.T."),
    ("LIM_FE", "Compras con factura electrónica (1%)", 240, "UVT", "Art. 336 num. 5 E.T."),
    ("SANC_MIN", "Sanción mínima (UVT del año de presentación)", 10, "UVT", "Art. 639 E.T."),
    ("TOPE_PAT", "Obligado a declarar por patrimonio", 4500, "UVT", "Art. 592 E.T."),
    ("TOPE_ING", "Obligado por ingresos / compras / consignaciones", 1400, "UVT", "Arts. 593 y 594-3 E.T."),
    ("TOPE_CONT", "Firma de contador público", 100000, "UVT", "Art. 596 E.T."),
    ("COMP_INF", "Componente inflacionario de rendimientos AG2025", 0.5543, "%", "Decreto 898 de 2026; arts. 38 y 40-1 E.T."),
    ("TAR_GO", "Tarifa ganancia ocasional general", 0.15, "%", "Art. 314 E.T."),
    ("TAR_LOT", "Tarifa loterías, rifas y apuestas", 0.20, "%", "Art. 317 E.T."),
    ("TAR_DIV2", "Tarifa previa dividendos 2.ª subcédula", 0.35, "%", "Art. 240 E.T."),
    ("REN_PRE", "Tarifa renta presuntiva", 0, "%", "Art. 188 E.T. (0% desde 2021)"),
    ("COMP_INF_GTO", "Componente inflacionario de costos financieros (no deducible)", 0.2835, "%", "Decreto 898 de 2026; arts. 81 y 118 E.T."),
    ("LIM_RAIS", "Cotización voluntaria RAIS (y 25% del ingreso)", 2500, "UVT", "Art. 55 E.T."),
]
for i, (cod, nom, val, uni, fte) in enumerate(params):
    r = 5 + i
    put(ws, f"A{r}", cod, S_note); put(ws, f"B{r}", nom, S_txt)
    fmt = "0.00%" if uni == "%" else "#,##0"
    put(ws, f"C{r}", val, S_in, fmt=fmt); put(ws, f"D{r}", uni, S_note); put(ws, f"E{r}", fte, S_note)
sec_row(ws, 29, "TABLA DEL ARTÍCULO 241 E.T. (cédulas general + pensiones + dividendos 1.ª subcédula)", "ABCDE")
header_row(ws, 30, "ABCDE", ["Tramo", "Desde UVT", "Hasta UVT", "Tarifa", "Suma fija UVT"])
for i, (d0, h0, t0, f0) in enumerate(P.tabla241):
    r = 31 + i
    put(ws, f"A{r}", i + 1, S_note); put(ws, f"B{r}", float(d0), S_in)
    put(ws, f"C{r}", float(h0), S_in); put(ws, f"D{r}", float(t0), S_in, fmt="0%")
    put(ws, f"E{r}", float(f0), S_in)
sec_row(ws, 40, "TABLA HISTÓRICA — DIVIDENDOS DE 2016 Y ANTERIORES (régimen de transición)", "ABCDE")
for i, (d0, h0, t0, f0) in enumerate(P.tabla_div2016):
    r = 41 + i
    put(ws, f"B{r}", float(d0), S_in); put(ws, f"C{r}", float(h0), S_in)
    put(ws, f"D{r}", float(t0), S_in, fmt="0%"); put(ws, f"E{r}", float(f0), S_in)
sec_row(ws, 46, "CESANTÍAS — PARTE NO GRAVADA SEGÚN SALARIO PROMEDIO 6 MESES (instructivo casilla 36, num. 4)", "ABCDE")
header_row(ws, 47, "BCD", ["Desde UVT", "Hasta UVT", "% no gravado"])
for i, (d0, h0, pc) in enumerate(P.tabla_cesantias):
    r = 48 + i
    put(ws, f"B{r}", float(d0), S_in); put(ws, f"C{r}", float(h0), S_in)
    put(ws, f"D{r}", float(pc), S_in, fmt="0%")
sec_row(ws, 56, "OTROS PARÁMETROS DE CONTROL", "ABCDE")
put(ws, "B57", "Beneficio de auditoría: impuesto neto mínimo del año anterior (71 UVT de 2024)", S_txt)
put(ws, "C57", "=ROUND(71*47065,0)", S_fx); put(ws, "E57", "Art. 689-3 E.T. (UVT 2024 = $47.065)", S_note)
put(ws, "B58", "GO exenta por venta de casa de habitación (con requisitos AFC)", S_txt)
put(ws, "C58", f"=ROUND(5000*{C5},-3)", S_fx); put(ws, "E58", "Art. 311-1 E.T. (5.000 UVT)", S_note)
put(ws, "H3", "PLAZOS 2026", S_sub); put(ws, "I3", "", S_sub); put(ws, "J3", "", S_sub)
header_row(ws, 4, "HIJ", ["Term.", "Fecha límite", "Fuente"])
cal = [(1,2,(8,12)),(3,4,(8,13)),(5,6,(8,14)),(7,8,(8,18)),(9,10,(8,19)),(11,12,(8,20)),(13,14,(8,21)),
       (15,16,(8,24)),(17,18,(8,25)),(19,20,(8,26)),(21,22,(8,27)),(23,24,(8,28)),(25,26,(8,31)),
       (27,28,(9,1)),(29,30,(9,2)),(31,32,(9,3)),(33,34,(9,4)),(35,36,(9,7)),(37,38,(9,8)),(39,40,(9,9)),
       (41,42,(9,10)),(43,44,(9,11)),(45,46,(9,14)),(47,48,(9,15)),(49,50,(9,16)),(51,52,(9,17)),
       (53,54,(9,18)),(55,56,(9,21)),(57,58,(9,22)),(59,60,(9,23)),(61,62,(9,24)),(63,64,(9,25)),
       (65,66,(9,28)),(67,68,(10,1)),(69,70,(10,2)),(71,72,(10,5)),(73,74,(10,6)),(75,76,(10,7)),
       (77,78,(10,8)),(79,80,(10,9)),(81,82,(10,13)),(83,84,(10,14)),(85,86,(10,15)),(87,88,(10,16)),
       (89,90,(10,19)),(91,92,(10,20)),(93,94,(10,21)),(95,96,(10,22)),(97,98,(10,23)),(99,0,(10,26))]
for dig in range(0, 100):
    if dig == 0: m, d0 = (10, 26)
    else:
        for a, b, (mm, dd) in cal:
            if dig in (a, b): m, d0 = (mm, dd); break
    rr = 5 + dig
    put(ws, f"H{rr}", dig, S_note); ws[f"H{rr}"].alignment = Alignment(horizontal="center")
    c = ws[f"I{rr}"]; c.value = dt.date(2026, m, d0); c.number_format = "dd/mm/yyyy"
    c.font = Font(name=F, size=8); c.alignment = Alignment(horizontal="center")
    put(ws, f"J{rr}", "Dcto. 2229/2023 art. 1.6.1.13.2.15", S_note)

# ============================================================ 03_PERFIL
ws = wb.create_sheet(SN["PERFIL"])
widths(ws, {"A": 44, "B": 18, "C": 16, "D": 30, "E": 48})
ws.sheet_view.showGridLines = False
titulo(ws, "ABCDE", f"PASO 1 · PERFIL DEL DECLARANTE — {CFG['nombre']}",
       "Complete solo celdas verdes. La identidad se valida contra el RUT, no contra este borrador.")
header_row(ws, 4, "ABCDE", ["Campo", "Valor", "Estado", "Fuente", "Explicación"])
for r, campo, val, est, fte, expl, fmt in CFG["PERFIL"]:
    if isinstance(val, str) and "{PAR}" in val:
        val = val.replace("{PAR}", PAR)
    put(ws, f"A{r}", campo, S_txt)
    if est == "CALCULADO":
        put(ws, f"B{r}", val, S_fx, fmt=fmt)
    else:
        put(ws, f"B{r}", val, S_in, fmt=fmt)
    stc = put(ws, f"C{r}", est, S_note)
    if est in ("FALTA", "POR CONFIRMAR", "PROVISIONAL"): S_warn(stc); stc.value = est
    put(ws, f"D{r}", fte, S_note); put(ws, f"E{r}", expl, S_note)
    ws.row_dimensions[r].height = 24

# ============================================================ 04_INGRESOS
ws = wb.create_sheet(SN["ING"])
widths(ws, {"A": 4, "B": 20, "C": 34, "D": 11, "E": 40, "F": 15, "G": 12, "H": 12, "I": 8, "J": 12, "K": 24, "L": 9, "M": 46})
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A5"
titulo(ws, "ABCDEFGHIJKLM", "PASO 2 · INGRESOS, COSTOS Y RETENCIONES — CONCILIACIÓN POR TERCERO",
       "Cada fila = un hecho económico. NO sume el certificado y la exógena: son dos fuentes del mismo ingreso. 'Incluir'=Sí lo lleva a la declaración. Valores en pesos EXACTOS (el redondeo del art. 577 se aplica solo en las casillas). ÁMBAR = estimado que usted debe reemplazar con certificado.")
header_row(ws, 4, "ABCDEFGHIJKLM", ["ID", "Categoría tributaria", "Tercero", "NIT", "Concepto",
    "Valor bruto exacto", "Costos", "Retención", "Incluir", "Soporte", "Fuente", "Casilla", "Observación"])
CATS = '"Trabajo - laboral,Trabajo - cesantías,Trabajo - sin costos,Trabajo - con costos,Capital - rendimientos,Capital - otros,No laboral,Pensión,Dividendo 1,Dividendo 2,Ganancia ocasional,Lotería"'
dv_cat = DataValidation(type="list", formula1=CATS, allow_blank=True); ws.add_data_validation(dv_cat)
dv_sn = DataValidation(type="list", formula1='"Sí,No"', allow_blank=False); ws.add_data_validation(dv_sn)
for i, rdata in enumerate(CFG["INGRESOS"]):
    r = 5 + i
    cat, ter, nit, con, bru, cos, ret, inc, sop, fte, casn, obs = rdata
    put(ws, f"A{r}", i + 1, S_note)
    put(ws, f"B{r}", cat, S_in, fmt="@")
    put(ws, f"C{r}", ter, S_txt, sz=8); put(ws, f"D{r}", nit, S_note)
    put(ws, f"E{r}", con, S_txt, sz=8)
    (S_est if sop == "ESTIMADO" else S_in)(put(ws, f"F{r}", float(bru)))
    put(ws, f"G{r}", float(cos), S_in); put(ws, f"H{r}", float(ret), S_in)
    put(ws, f"I{r}", inc, S_in, fmt="@")
    sc = put(ws, f"J{r}", sop, S_note)
    if sop in ("Pendiente", "Provisional", "ESTIMADO"): S_warn(sc); sc.value = sop
    put(ws, f"K{r}", fte, S_note); put(ws, f"L{r}", casn, S_note); put(ws, f"M{r}", obs, S_note)
    ws.row_dimensions[r].height = 26
dv_cat.add("B5:B24"); dv_sn.add("I5:I24")
SS = lambda cat, campo="F": f'SUMIFS(${campo}$5:${campo}$24,$B$5:$B$24,"{cat}",$I$5:$I$24,"Sí")'
sec_row(ws, 26, "RESUMEN CALCULADO — VALORES EXACTOS (antes del redondeo del formulario)", "ABCDEF")
res = [
    (27, "Total ingresos brutos incluidos", '=SUMIFS($F$5:$F$24,$I$5:$I$24,"Sí")'),
    (28, "Rentas de trabajo — casilla 32 (laboral + cesantías + sin costos)",
     f'={SS("Trabajo - laboral")}+{SS("Trabajo - cesantías")}+{SS("Trabajo - sin costos")}'),
    (29, "Trabajo con costos — casilla 43", f'={SS("Trabajo - con costos")}'),
    (30, "Rendimientos financieros elegibles (comp. inflacionario)", f'={SS("Capital - rendimientos")}'),
    (31, "Otros ingresos de capital", f'={SS("Capital - otros")}'),
    (32, "Rentas no laborales", f'={SS("No laboral")}'),
    (33, "Pensiones", f'={SS("Pensión")}'),
    (34, "Dividendos 1.ª subcédula", f'={SS("Dividendo 1")}'),
    (35, "Dividendos 2.ª subcédula", f'={SS("Dividendo 2")}'),
    (36, "Ganancias ocasionales (sin loterías)", f'={SS("Ganancia ocasional")}'),
    (37, "Loterías, rifas y apuestas", f'={SS("Lotería")}'),
    (38, "Retenciones incluidas", '=SUMIFS($H$5:$H$24,$I$5:$I$24,"Sí")'),
    (39, "Costos — trabajo con costos", f'={SS("Trabajo - con costos", "G")}'),
    (40, "Costos financieros — capital", f'={SS("Capital - rendimientos", "G")}+{SS("Capital - otros", "G")}'),
    (41, "Costos — no laborales", f'={SS("No laboral", "G")}'),
    (42, "Solo 'Trabajo - laboral' (base 10% dependientes art. 387)", f'={SS("Trabajo - laboral")}'),
    (43, "Partidas incluidas sin soporte confirmado",
     '=COUNTIFS($I$5:$I$24,"Sí",$J$5:$J$24,"Provisional")+COUNTIFS($I$5:$I$24,"Sí",$J$5:$J$24,"Pendiente")+COUNTIFS($I$5:$I$24,"Sí",$J$5:$J$24,"ESTIMADO")'),
    (44, "Cesantías incluidas", f'={SS("Trabajo - cesantías")}'),
]
for r, lab, fx in res:
    put(ws, f"B{r}", lab, S_txt); put(ws, f"F{r}", fx, S_fx)
put(ws, "B45", "Parte NO gravada de cesantías (% según salario promedio, tabla instructivo c.36)", S_txt)
put(ws, "F45", f"=VLOOKUP(MAX(0,{PER}!$B$28-0.000000001),{PAR}!$B$48:$D$54,3,TRUE)", S_fx, fmt="0%")
put(ws, "B46", "Cesantías exentas (num. 4 art. 206)", S_txt)
put(ws, "F46", "=MIN(F44,F44*F45)", S_fx)

# ============================================================ 05_BENEFICIOS
ws = wb.create_sheet(SN["BEN"])
widths(ws, {"A": 5, "B": 44, "C": 14, "D": 9, "E": 8, "F": 13, "G": 8, "H": 15, "I": 42, "J": 46})
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A5"
titulo(ws, "ABCDEFGHIJ", "PASO 3 · BENEFICIOS: INCRNGO, EXENTAS, DEDUCCIONES Y DESCUENTOS",
       "La columna F calcula el valor ACEPTADO por la regla legal. 'Aplicar'=Sí exige soporte. ÁMBAR = valor ESTIMADO que usted debe reemplazar.")
header_row(ws, 4, "ABCDEFGHIJ", ["ID", "Concepto", "Base / entrada", "Meses/%", "Aplicar", "Valor aceptado", "Destino", "Soporte", "Regla resumida", "Advertencia"])
dv_sn2 = DataValidation(type="list", formula1='"Sí,No"', allow_blank=False); ws.add_data_validation(dv_sn2)
for (i, nom, base, dd, apl, dest, sop, regla, adv) in BENS:
    r = i
    put(ws, f"A{r}", i, S_note); put(ws, f"B{r}", nom, S_txt, sz=8)
    est_fill = sop == "ESTIMADO"
    if i == 10:
        put(ws, f"C{r}", f"={ING}!F30", S_fx); put(ws, f"D{r}", f"={PAR}!$C$21", S_fx, fmt="0.00%")
        put(ws, f"E{r}", "Sí", S_in, fmt="@")
        put(ws, f"F{r}", f'=IF(AND(E10="Sí",{PER}!$B$17="No"),C10*D10,0)', S_fx)
    elif i == 13:
        put(ws, f"C{r}", '=IF(D13="220",C11,C12)', S_fx)
        put(ws, f"D{r}", dd if isinstance(dd, str) and dd in ("FNA", "220") else "FNA", S_in, fmt="@")
        put(ws, f"E{r}", apl, S_in, fmt="@")
        put(ws, f"F{r}", '=IF(E13="Sí",C13,0)', S_fx)
    elif i == 21:
        put(ws, f"C{r}", f"={DEP}!B20", S_fx, fmt="0"); put(ws, f"D{r}", f"={DEP}!B19", S_fx, fmt="0")
        put(ws, f"E{r}", '=IF(C21>0,"Sí","No")', S_fx, fmt="@")
        put(ws, f"F{r}", f'=IF(C21>0,MIN(0.10*{ING}!$F$42,D21*{PAR}!$C$13*{PAR}!$C$5),0)', S_fx)
    elif i == 22:
        put(ws, f"C{r}", f"={DEP}!B18", S_fx, fmt="0"); put(ws, f"D{r}", f"={PAR}!$C$14", S_fx, fmt="0")
        put(ws, f"E{r}", '=IF(C22>0,"Sí","No")', S_fx, fmt="@")
        put(ws, f"F{r}", f'=IF(C22>0,MIN(C22,{PAR}!$C$15)*{PAR}!$C$14*{PAR}!$C$5,0)', S_fx)
    else:
        (S_est if est_fill else S_in)(put(ws, f"C{r}", float(base) if base is not None else 0))
        put(ws, f"D{r}", float(dd) if not isinstance(dd, str) else dd, S_in, fmt="@" if isinstance(dd, str) else "0.00%")
        if isinstance(dd, (int, float, Decimal)) and float(dd or 0) in (0.0, 12.0):
            ws[f"D{r}"].number_format = "0"
        put(ws, f"E{r}", apl, S_in, fmt="@")
        fmap = {
            7: f'=IF(E7="Sí",MIN(C7,0.25*{ING}!$F$27,{PAR}!$C$27*{C5}),0)',
            17: f'=IF(E17="Sí",MIN(C17,{PAR}!$C$10*{C5}),0)',
            18: f'=IF(E18="Sí",MIN(C18,{PAR}!$C$11*{C5}),0)',
            19: f'=IF(E19="Sí",MIN(C19,D19*{PAR}!$C$12*{C5}),0)',
            20: '=IF(E20="Sí",C20*D20,0)',
            23: f'=IF(E23="Sí",MIN(C23*D23,{PAR}!$C$16*{C5}),0)',
            37: '=IF(E37="Sí",0.25*C37,0)',
        }
        put(ws, f"F{r}", fmap.get(i, f'=IF(E{r}="Sí",C{r},0)'), S_fx)
    put(ws, f"G{r}", dest, S_note)
    sc = put(ws, f"H{r}", sop, S_note)
    if sop in ("Pendiente", "Provisional", "ESTIMADO", "En conflicto", "Falta", "Pendiente soporte"):
        S_warn(sc); sc.value = sop
    put(ws, f"I{r}", regla, S_note); put(ws, f"J{r}", adv, S_note)
    ws.row_dimensions[r].height = 24
for r in list(range(5, 10)) + list(range(14, 21)) + list(range(23, 41)):
    dv_sn2.add(f"E{r}")
dv_afc = DataValidation(type="list", formula1='"FNA,220"'); ws.add_data_validation(dv_afc); dv_afc.add("D13")
sec_row(ws, 43, "CASOS ESPECIALES POR CASILLA — use solo si aplican y existe soporte", "ABCDEFGHIJ")
header_row(ws, 44, "ABCDEFGH", ["ID", "Concepto", "Valor", "Aplicar", "Aceptado", "Destino", "Soporte", "Advertencia"])
ESPECIALES = [
    (45, "Otras rentas exentas — trabajo con costos", 0, "No", "48", "Sin 25% si optó por costos"),
    (46, "Intereses de vivienda — trabajo con costos", 0, "No", "50", "No duplicar certificado"),
    (47, "Otras deducciones — trabajo con costos", 0, "No", "51", ""),
    (48, "Pérdidas compensables — trabajo con costos", 0, "No", "56", "Misma subcédula"),
    (49, "Rentas líquidas pasivas ECE — capital", 0, "No", "62", "Requiere análisis ECE (contador)"),
    (50, "Otras rentas exentas — capital", 0, "No", "64", ""),
    (51, "Intereses de vivienda — capital", 0, "No", "66", "No duplicar certificado"),
    (52, "Otras deducciones — capital", 0, "No", "67", ""),
    (53, "Pérdidas compensables — capital", 0, "No", "72", "Misma subcédula"),
    (54, "Devoluciones, rebajas y descuentos — no laboral", 0, "No", "75", ""),
    (55, "INCRNGO — no laboral", 0, "No", "76", ""),
    (56, "Rentas líquidas pasivas ECE — no laboral", 0, "No", "79", "Requiere análisis ECE"),
    (57, "Otras rentas exentas — no laboral", 0, "No", "81", ""),
    (58, "Intereses de vivienda — no laboral", 0, "No", "83", "No duplicar certificado"),
    (59, "Otras deducciones — no laboral", 0, "No", "84", ""),
    (60, "Pérdidas compensables — no laboral", 0, "No", "89", "Misma subcédula"),
    (61, "INCRNGO dividendos 2016 y anteriores", 0, "No", "105", ""),
    (62, "Dividendos recibidos del exterior", 0, "No", "109", "Requiere contador"),
    (63, "Renta exenta sobre dividendos del exterior", 0, "No", "110", "Solo CAN probada"),
    (64, "Rentas gravables de activos exceptuados de presuntiva", 0, "No", "98", "Instructivo c.98"),
    (65, "Rentas exentas laborales FUERA del límite del 40% (206 n.6-9, 206-1, CAN)", 0, "No", "36/41",
     "Solo con norma y soporte expresos"),
]
for (i, nom, val, apl, dest, adv) in ESPECIALES:
    r = i
    adv = CFG["ESPECIALES_NOTAS"].get(i, adv)
    put(ws, f"A{r}", f"E-{i-44:02d}", S_note); put(ws, f"B{r}", nom, S_txt, sz=8)
    put(ws, f"C{r}", float(val), S_in); put(ws, f"D{r}", apl, S_in, fmt="@")
    put(ws, f"E{r}", f'=IF(D{r}="Sí",C{r},0)', S_fx)
    put(ws, f"F{r}", dest, S_note); put(ws, f"G{r}", "Pendiente", S_note)
    nc = put(ws, f"H{r}", adv, S_note)
    if i in CFG["ESPECIALES_NOTAS"]: S_warn(nc); nc.value = adv
    dv_sn2.add(f"D{r}")
    ws.row_dimensions[r].height = 18 if i not in CFG["ESPECIALES_NOTAS"] else 34

# ============================================================ 06_DEPENDIENTES
ws = wb.create_sheet(SN["DEP"])
widths(ws, {"A": 9, "B": 15, "C": 9, "D": 9, "E": 9, "F": 12, "G": 12, "H": 60, "I": 8, "J": 30})
ws.sheet_view.showGridLines = False
titulo(ws, "ABCDEFGHIJ", "PASO 4 · DEPENDIENTES — REGISTRO ÚNICO Y CONTROL DE CONCURRENCIA (DUR 1.2.1.20.3)",
       "Una fila = una persona. El mismo dependiente puede dar las DOS deducciones (art. 387 y 72 UVT) solo si usted tiene rentas de una relación laboral o legal y reglamentaria.")
header_row(ws, 4, "ABCDEFGHIJ", ["Alias", "Relación", "Meses", "Art. 387", "72 UVT", "¿Renta laboral?", "Soporte", "Identificación y observación", "Cuenta", "Control"])
dv_sn3 = DataValidation(type="list", formula1='"Sí,No"'); ws.add_data_validation(dv_sn3)
for i, (ali, rel, mes, a387, u72, lab, sop, obs) in enumerate(CFG["DEPENDIENTES"]):
    r = 5 + i
    put(ws, f"A{r}", ali, S_note); put(ws, f"B{r}", rel, S_in, fmt="@")
    put(ws, f"C{r}", mes, S_in, fmt="0"); put(ws, f"D{r}", a387, S_in, fmt="@")
    put(ws, f"E{r}", u72, S_in, fmt="@"); put(ws, f"F{r}", lab, S_in, fmt="@")
    sc = put(ws, f"G{r}", sop, S_note)
    if sop == "Pendiente": S_warn(sc); sc.value = sop
    put(ws, f"H{r}", obs, S_note)
    put(ws, f"I{r}", f'=IF(AND(OR(D{r}="Sí",E{r}="Sí"),F{r}="Sí"),1,0)', S_fx, fmt="0")
    put(ws, f"J{r}", f'=IF(AND(OR(D{r}="Sí",E{r}="Sí"),F{r}<>"Sí"),"BLOQUEO: exige renta laboral",IF(AND(OR(D{r}="Sí",E{r}="Sí"),G{r}<>"Disponible"),"PENDIENTE SOPORTE","OK"))', S_fx, fmt="@")
    for col in "DEF": dv_sn3.add(f"{col}{r}")
    ws.row_dimensions[r].height = 24
sec_row(ws, 15, "RESUMEN CALCULADO", "ABCD")
put(ws, "A16", "Concepto", S_sub); put(ws, "B16", "Valor", S_sub); put(ws, "C16", "Uso", S_sub)
put(ws, "A17", "Dependientes únicos activos (casilla 138)", S_txt); put(ws, "B17", "=SUM(I5:I12)", S_fx, fmt="0"); put(ws, "C17", "Casilla 138", S_note)
put(ws, "A18", "Dependientes válidos para 72 UVT (máx. 4)", S_txt)
put(ws, "B18", f'=MIN({PAR}!$C$15,SUMPRODUCT((E5:E12="Sí")*(F5:F12="Sí")))', S_fx, fmt="0"); put(ws, "C18", "Casilla 139", S_note)
put(ws, "A19", "Meses cubiertos para art. 387 (máximo entre activos)", S_txt)
put(ws, "B19", '=MAX(IF(AND(D5="Sí",F5="Sí"),C5,0),IF(AND(D6="Sí",F6="Sí"),C6,0),IF(AND(D7="Sí",F7="Sí"),C7,0),IF(AND(D8="Sí",F8="Sí"),C8,0),IF(AND(D9="Sí",F9="Sí"),C9,0),IF(AND(D10="Sí",F10="Sí"),C10,0),IF(AND(D11="Sí",F11="Sí"),C11,0),IF(AND(D12="Sí",F12="Sí"),C12,0))', S_fx, fmt="0")
put(ws, "C19", "Deducción mensual 32 UVT", S_note)
put(ws, "A20", "Dependientes activos art. 387", S_txt)
put(ws, "B20", '=SUMPRODUCT((D5:D12="Sí")*(F5:F12="Sí"))', S_fx, fmt="0")

# ============================================================ 07_PATRIMONIO (con columna histórica)
ws = wb.create_sheet(SN["PAT"])
widths(ws, {"A": 4, "B": 9, "C": 16, "D": 42, "E": 15, "F": 10, "G": 15, "H": 15, "I": 8, "J": 13, "K": 26, "L": 14, "M": 14, "N": 52})
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A5"
titulo(ws, "ABCDEFGHIJKLMN", "PASO 5 · PATRIMONIO Y DEUDAS A 31 DE DICIEMBRE DE 2025 — CON HISTORIA 2024",
       "Activos: se declara el MAYOR entre (valor reportado × participación) y el 'Alterno' (costo fiscal o avalúo, art. 277). Deudas: su parte real con soporte. Columnas L-M: lo declarado en AG2024 y la variación (la DIAN también la ve). ÁMBAR = estimado a reemplazar.")
header_row(ws, 4, "ABCDEFGHIJKLMN", ["ID", "Tipo", "Categoría", "Entidad / activo", "Valor reportado", "Partic.",
    "Alterno (costo/avalúo)", "Valor reconocido", "Incluir", "Soporte", "Fuente", "Declarado AG2024", "Variación", "Observación"])
PAT_ROWS = list(CFG["PATRIMONIO"])
while len(PAT_ROWS) < 20:
    PAT_ROWS.append(("Activo", "Otro", "Fila libre", 0, 1, 0, "No", "Pendiente", "", 0, ""))
dv_td = DataValidation(type="list", formula1='"Activo,Deuda"'); ws.add_data_validation(dv_td)
for i, (tipo, catg, ent, rep, par, alt, inc, sop, fte, hist, obs) in enumerate(PAT_ROWS):
    r = 5 + i
    put(ws, f"A{r}", i + 1, S_note); put(ws, f"B{r}", tipo, S_in, fmt="@")
    put(ws, f"C{r}", catg, S_note); put(ws, f"D{r}", ent, S_txt, sz=8)
    (S_est if sop == "ESTIMADO" else S_in)(put(ws, f"E{r}", float(rep)))
    put(ws, f"F{r}", float(par), S_in, fmt="0.00%")
    put(ws, f"G{r}", float(alt), S_in)
    put(ws, f"H{r}", f'=IF(I{r}<>"Sí",0,IF(B{r}="Deuda",IF(G{r}=0,E{r}*F{r},G{r}),MAX(E{r}*F{r},G{r})))', S_fx)
    put(ws, f"I{r}", inc, S_in, fmt="@"); dv_sn3b = None
    sc = put(ws, f"J{r}", sop, S_note)
    if sop not in ("", "Confirmado"):
        (S_alert if sop in ("CRÍTICO",) else S_warn)(sc); sc.value = sop
    put(ws, f"K{r}", fte, S_note)
    hc = ws[f"L{r}"]; hc.value = float(hist); hc.number_format = "#,##0"
    hc.font = Font(name=F, size=8, color=GRIS); hc.alignment = Alignment(horizontal="right", vertical="center")
    put(ws, f"M{r}", f'=IF(I{r}="Sí",H{r},0)-L{r}', S_fx, fmt="#,##0;[Red]-#,##0")
    ws[f"M{r}"].font = Font(name=F, size=8, color=GRIS)
    put(ws, f"N{r}", obs, S_note)
    ws.row_dimensions[r].height = 26
dv_td.add("B5:B24")
dv_iy = DataValidation(type="list", formula1='"Sí,No"'); ws.add_data_validation(dv_iy); dv_iy.add("I5:I24")
sec_row(ws, 26, "RESUMEN Y CONCILIACIÓN CONTRA AG2024", "ABCDEFGH")
put(ws, "D27", "Opción", S_sub); put(ws, "E27", "Valor", S_sub)
patres = [
    (28, "Activos identificados e incluidos", "Automático", f'=SUMIFS($H$5:$H$24,$B$5:$B$24,"Activo",$I$5:$I$24,"Sí")'),
    (29, "Diferencia frente al patrimonio bruto AG2024 (solo control)", "Control", f"=E28-{PER}!B22"),
    (30, "¿Sumar el faltante frente a AG2024 como ajuste provisional? (NO recomendado)", "No", None),
    (31, "PATRIMONIO BRUTO usado en el Formulario 210", "Calculado", f'=E28+IF(D30="Sí",MAX(0,{PER}!B22-E28),0)'),
    (32, "Deudas identificadas e incluidas", "Calculado", f'=SUMIFS($H$5:$H$24,$B$5:$B$24,"Deuda",$I$5:$I$24,"Sí")'),
    (33, "Patrimonio líquido", "Calculado", "=MAX(0,E31-E32)"),
    (34, "Variación del patrimonio BRUTO vs. AG2024", "Control", f"=E31-{PER}!B22"),
    (35, "Variación del patrimonio LÍQUIDO vs. AG2024", "Control", f"=E33-{PER}!B24"),
]
for r, lab, opc, fx in patres:
    put(ws, f"A{r}", lab, S_txt); ws.merge_cells(f"A{r}:C{r}")
    if r == 30:
        put(ws, f"D{r}", "No", S_in, fmt="@"); dv_iy.add(f"D{r}")
    else:
        put(ws, f"D{r}", opc, S_note)
    if fx: put(ws, f"E{r}", fx, S_fx, tot=(r in (31, 32, 33)))
    if fx: ws[f"E{r}"].number_format = "#,##0;[Red]-#,##0"
pat_warn = CFG.get("PAT_WARN", "La variación del patrimonio líquido debe caber en lo que la renta explica (arts. 236-239): el control de 12_CONTROLES lo verifica. Revise la columna Variación fila por fila y documente TODO salto material antes de presentar.")
put(ws, "A37", pat_warn, S_warn)
ws.merge_cells("A37:N39"); ws.row_dimensions[37].height = 22

# ============================================================ 08_HISTORICO
ws = wb.create_sheet(SN["HIST"])
widths(ws, {"A": 16, "B": 42, "C": 14, "D": 14, "E": 9, "F": 15, "G": 14, "H": 14, "I": 56})
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A5"
titulo(ws, "ABCDEFGHI", "REGISTRO HISTÓRICO — DE DÓNDE VIENE CADA CIFRA (AG2023 → AG2024 → AG2025)",
       "Reconstrucción de su AyudaRenta 2024 presentada, partida por partida, conectada EN VIVO con este archivo (columna 'AG2025 (este archivo)'). Así cada año arranca del anterior sin perder el hilo. Gris = dato histórico fijo; azul = valor vivo de 2025.")
header_row(ws, 4, "ABCDEFGHI", ["Sección", "Entidad / concepto", "AG2023", "AG2024", "¿2025?", "Dónde está ahora",
                                "AG2025 (este archivo)", "Variación 24→25", "Nota"])

def hist_ref(kind, x):
    if kind == "PAT": return f"={PAT}!H{x}"
    if kind == "ING": return f"={ING}!F{x}"
    if kind == "BEN": return f"={BEN}!F{x}"
    if kind == "F210": return f"={F210}!D{fila_casilla(x)}"
    return None

HISTR = CFG.get("HISTORICO") or [
    ("PATRIMONIO", None, None, None, None, None, None),
    ("Registro", "Sin datos históricos cargados en este caso", None, None, "", None,
     "Llene CASO['HISTORICO'] con la composición de las declaraciones anteriores (esquema en caso_demo.py) para que cada cifra tenga trazabilidad año a año."),
]
r = 5
for row in HISTR:
    sec, ent, v23, v24, sigue, ref, nota = row
    if ent is None:
        sec_row(ws, r, sec, "ABCDEFGHI", S_sec2); r += 1; continue
    put(ws, f"A{r}", sec, S_note); put(ws, f"B{r}", ent, S_txt, sz=8)
    for col, v in (("C", v23), ("D", v24)):
        c = ws[f"{col}{r}"]
        c.value = (float(v) if v is not None else None)
        c.number_format = "#,##0"; c.font = Font(name=F, size=8, color=GRIS)
        c.alignment = Alignment(horizontal="right", vertical="center")
        if v is None: c.value = "n.d."; c.number_format = "@"
    put(ws, f"E{r}", sigue, S_note)
    if ref:
        put(ws, f"F{r}", f"{SN[ref[0]].split('_',1)[0]} · {ref[0]}", S_note)
        put(ws, f"G{r}", hist_ref(*ref), S_fx)
        put(ws, f"H{r}", f'=IF(ISNUMBER(D{r}),G{r}-D{r},"")', S_fx, fmt="#,##0;[Red]-#,##0")
        ws[f"H{r}"].font = Font(name=F, size=8, color=GRIS)
    else:
        put(ws, f"F{r}", "—", S_note); put(ws, f"G{r}", "—", S_note)
    put(ws, f"I{r}", nota, S_note)
    ws.row_dimensions[r].height = 24
    r += 1
put(ws, f"A{r+1}", "CÓMO USARLO EL PRÓXIMO AÑO: al preparar la siguiente declaración, la columna del año actual se convierte en la nueva columna histórica. Así ninguna cifra vuelve a aparecer de la nada — el origen típico de los errores heredados.", S_warn)
ws.merge_cells(f"A{r+1}:I{r+2}")

print("hojas 03-08 listas")

# ============================================================ 09/10 FORMULARIO_210 y MOTOR_EXACTO
ESP = lambda i: f"{BEN}!$E${i}"
BF = lambda i: f"{BEN}!$F${i}"
FD = lambda c: f"{F210}!D{fila_casilla(c)}"

def tarifa_fx(base_cell, tabla_rng, oficial):
    b = f"TRUNC({base_cell}/{PAR}!$C$5,2)"
    lk = f"MAX(0,{b}-0.000000001)"
    low = f"VLOOKUP({lk},{tabla_rng},1,TRUE)"
    rate = f"VLOOKUP({lk},{tabla_rng},3,TRUE)"
    fija = f"VLOOKUP({lk},{tabla_rng},4,TRUE)"
    core = f"(TRUNC(({b}-{low})*{rate}+0.000000001,2)+{fija})*{PAR}!$C$5"
    return f"ROUND({core},-3)" if oficial else core

T241 = f"{PAR}!$B$31:$E$37"
T2016 = f"{PAR}!$B$41:$E$44"

def spec(o):
    R = (lambda x: f"ROUND({x},-3)") if o else (lambda x: f"({x})")
    tope_afc = (f"MIN({BEN}!$F$13+{BEN}!$F$14+{BEN}!$F$15+{BEN}!$F$24+{BEN}!$F$25+{BEN}!$F$26,"
                f"0.30*{ING}!$F$27,{PAR}!$C$9*{PAR}!$C$5)")
    e65 = ESP(65)
    s = {}
    def C(cas, formula, concepto, simple="", norma=""):
        s[fila_casilla(cas)] = (str(cas), concepto, formula, simple, norma)
    C(28, f'=IF(D22+D39+D57+D76>0,{R(BF(23))},0)', "Uno por ciento (1%) de compras con factura electrónica",
      "1% de sus compras con factura electrónica (tope 240 UVT). FUERA del límite del 40% y solo si hay renta líquida (no genera pérdida).", "Art. 336 num. 5")
    C(29, f'={R(f"{PAT}!$E$31")}', "Total patrimonio bruto", "Todo lo que posee a 31/12/2025 (hoja 07_PATRIMONIO).", "Arts. 261-277")
    C(30, f'={R(f"{PAT}!$E$32")}', "Deudas", "Su parte real de las deudas con soporte.", "Arts. 283-284, 767, 770")
    C(31, "=MAX(0,D7-D8)", "Total patrimonio líquido (29-30)", "", "Art. 282")
    C(32, f'={R(f"{ING}!$F$28")}', "Ingresos brutos por rentas de trabajo",
      "Salarios + prestaciones + otros pagos + cesantías + servicios personales SIN costos (instructivo c.32).", "Art. 103; instructivo c.32")
    C(33, f'={R(f"{BEN}!$F$5+{BEN}!$F$6+{BEN}!$F$7+{BEN}!$F$8")}', "Ingresos no constitutivos de renta",
      "Aportes obligatorios a salud y pensión (y RAIS voluntario si aplica).", "Arts. 55-56")
    C(34, "=MAX(0,D12-D13)", "Renta líquida (32-33)", "", "Art. 26")
    C(35, f'={R(f"MIN({BEN}!$F$13+{BEN}!$F$14+{BEN}!$F$15,{tope_afc})")}',
      "Rentas exentas — aportes AFC, FVP y AVC", "Tope conjunto: 30% del ingreso del año y 3.800 UVT (se reparte entre subcédulas).", "Arts. 126-1 y 126-4")
    ces_ex = R(f"{ING}!$F$46"); ces_tot = R(f"{ING}!$F$44")
    C(36, f'={R(f"MIN(D14-D15,MAX(0,{BF(16)}+{e65}+{ces_ex})+MIN({PAR}!$C$8*{PAR}!$C$5,0.25*MAX(0,D14-D15-D18-D19-{BF(16)}-{e65}-{ces_tot})))")}',
      "Otras rentas exentas (incluye el 25% laboral)",
      "25% de los pagos laborales tras restar INCRNGO, AFC, deducciones y demás exentas; tope 790 UVT. NO se aplica sobre cesantías (num. 4 propio).",
      "Art. 206 num. 10; instructivo c.36")
    C(37, "=D15+D16", "Total rentas exentas de trabajo (35+36)", "", "")
    C(38, f'={R(f"MIN(D14,{BF(17)})")}', "Deducción — intereses de vivienda", "Tope anual conjunto de 1.200 UVT entre subcédulas.", "Art. 119")
    C(39, f'={R(f"MIN(D14,{BF(18)}+{BF(19)}+{BF(20)}+{BF(21)})")}', "Deducción — otras deducciones imputables",
      "ICETEX + medicina prepagada + 50% GMF + dependientes art. 387 (mensualizado).", "Arts. 115, 119, 387")
    C(40, "=D18+D19", "Total deducciones imputables (38+39)", "", "")
    C(41, f'={R(f"MIN(D14,{e65}+MIN(MAX(0,D14-{e65}),MAX(0,D17+D20-{e65}),D80))")}',
      "Rentas exentas y deducciones imputables (LIMITADAS)",
      "Primer uso del límite del 40%/1.340 UVT; lo 'fuera del límite' (E-21) pasa sin consumir cupo.", "Art. 336 num. 3")
    C(42, "=MAX(0,D14-D21)", "Renta líquida ordinaria de trabajo (34-41)", "", "")
    C(43, f'={R(f"{ING}!$F$29")}', "Ingresos brutos trabajo con costos", "Honorarios/servicios que OPTAN por costos (renuncian al 25%).", "Art. 336 num. 4")
    C(44, f'={R(f"MIN({BF(9)},{ING}!$F$29)")}', "INCRNGO de trabajo con costos", "", "")
    C(45, f'={R(f"{ING}!$F$39")}', "Costos y gastos procedentes", "", "Art. 107")
    C(46, "=MAX(0,D25-D26-D27)", "Renta líquida (43-44-45)", "", "")
    C(47, f'={R(f"MIN({BF(24)},MAX(0,{tope_afc}-D15))")}', "AFC, FVP y AVC", "Segundo uso del tope AFC.", "")
    C(48, f'={R(ESP(45))}', "Otras rentas exentas", "", "")
    C(49, "=D29+D30", "Total rentas exentas (47+48)", "", "")
    C(50, f'={R(f"MIN(D28,{ESP(46)},MAX(0,{PAR}!$C$10*{PAR}!$C$5-D18))")}', "Intereses de vivienda", "Saldo del tope de 1.200 UVT.", "")
    C(51, f'={R(f"MIN(D28,{ESP(47)})")}', "Otras deducciones imputables", "", "")
    C(52, "=D32+D33", "Total deducciones (50+51)", "", "")
    C(53, f'={R(f"MIN(D28,D31+D34,MAX(0,D80-MAX(0,D21-{e65})))")}', "Exentas y deducciones LIMITADAS", "Segundo uso del límite del 40%.", "")
    C(54, "=MAX(0,D25-D26-D27-D35)", "Renta líquida ordinaria (43-44-45-53)", "", "")
    C(55, "=MAX(0,D26+D27-D25)", "Pérdida líquida del ejercicio", "", "")
    C(56, f'={R(f"MIN(D36,{ESP(48)})")}', "Compensaciones por pérdidas", "", "")
    C(57, "=MAX(0,D36-D38)", "Renta líquida (54-56)", "", "")
    C(58, f'={R(f"{ING}!$F$30+{ING}!$F$31")}', "Ingresos brutos rentas de capital", "Intereses, rendimientos, arrendamientos.", "Art. 335")
    C(59, f'={R(BF(10))}', "INCRNGO (componente inflacionario)",
      "55,43% de los rendimientos BRUTOS de entidades vigiladas (no obligados a contabilidad).", "Arts. 38 y 40-1; Dcto. 898/2026")
    q = '"'
    C(60, f'={R(f"IF({PER}!$B$17=" + q + "No" + q + f",{ING}!$F$40*(1-{PAR}!$C$26),{ING}!$F$40)")}',
      "Costos y gastos procedentes", "Si no lleva contabilidad, el 28,35% de los costos financieros NO es deducible.", "Arts. 81 y 118; Dcto. 898/2026")
    C(61, "=MAX(0,D42-D43-D44)", "Renta líquida (58-59-60)", "", "")
    C(62, f'={R(ESP(49))}', "Rentas líquidas pasivas — ECE", "Solo régimen ECE (requiere contador).", "Arts. 882 y ss.")
    C(63, f'={R(f"MIN({BF(25)},MAX(0,{tope_afc}-D15-D29))")}', "AFC, FVP y AVC", "", "")
    C(64, f'={R(ESP(50))}', "Otras rentas exentas", "", "")
    C(65, "=D47+D48", "Total rentas exentas (63+64)", "", "")
    C(66, f'={R(f"MIN(D45+D46,{ESP(51)},MAX(0,{PAR}!$C$10*{PAR}!$C$5-D18-D32))")}', "Intereses de vivienda", "", "")
    C(67, f'={R(f"MIN(D45+D46,{ESP(52)})")}', "Otras deducciones imputables", "", "")
    C(68, "=D50+D51", "Total deducciones (66+67)", "", "")
    C(69, f'={R(f"MIN(D45+D46,D49+D52,MAX(0,D80-MAX(0,D21-{e65})-D35))")}', "Exentas y deducciones LIMITADAS", "Tercer uso del límite.", "")
    C(70, "=MAX(0,D42+D46-D43-D44-D53)", "Renta líquida ordinaria (58+62-59-60-69)", "", "")
    C(71, "=MAX(0,D43+D44-D42-D46)", "Pérdida líquida del ejercicio", "", "")
    C(72, f'={R(f"MIN(D54,{ESP(53)})")}', "Compensaciones por pérdidas", "", "")
    C(73, "=MAX(0,D54-D56)", "Renta líquida de capital (70-72)", "", "")
    C(74, f'={R(f"{ING}!$F$32")}', "Ingresos brutos rentas no laborales", "", "Art. 335")
    C(75, f'={R(f"MIN(D60,{ESP(54)})")}', "Devoluciones, rebajas y descuentos", "", "")
    C(76, f'={R(f"MIN(MAX(0,D60-D61),{ESP(55)})")}', "Ingresos no constitutivos de renta", "", "")
    C(77, f'={R(f"{ING}!$F$41")}', "Costos y gastos procedentes", "", "")
    C(78, "=MAX(0,D60-D61-D62-D63)", "Renta líquida (74-75-76-77)", "", "")
    C(79, f'={R(ESP(56))}', "Rentas líquidas pasivas — ECE", "", "")
    C(80, f'={R(f"MIN({BF(26)},MAX(0,{tope_afc}-D15-D29-D47))")}', "AFC, FVP y AVC", "", "")
    C(81, f'={R(ESP(57))}', "Otras rentas exentas", "", "")
    C(82, "=D66+D67", "Total rentas exentas (80+81)", "", "")
    C(83, f'={R(f"MIN(D64+D65,{ESP(58)},MAX(0,{PAR}!$C$10*{PAR}!$C$5-D18-D32-D50))")}', "Intereses de vivienda", "", "")
    C(84, f'={R(f"MIN(D64+D65,{ESP(59)})")}', "Otras deducciones imputables", "", "")
    C(85, "=D69+D70", "Total deducciones (83+84)", "", "")
    C(86, f'={R(f"MIN(D64+D65,D68+D71,MAX(0,D80-MAX(0,D21-{e65})-D35-D53))")}', "Exentas y deducciones LIMITADAS", "Cuarto uso del límite.", "")
    C(87, "=MAX(0,D60-D61-D62-D63+D65-D72)", "Renta líquida ordinaria", "", "")
    C(88, "=MAX(0,D61+D62+D63-D60-D65)", "Pérdida líquida del ejercicio", "", "")
    C(89, f'={R(f"MIN(D73,{ESP(60)})")}', "Compensaciones por pérdidas", "", "")
    C(90, "=MAX(0,D73-D75)", "Renta líquida no laboral (87-89)", "", "")
    C(91, "=MAX(0,D21+D22+D35+D39+D53+D57+D72+D76)", "Renta líquida cédula general (41+42+53+57+69+73+86+90)",
      "Suma de subcédulas antes de restar los beneficios de la casilla 92.", "Art. 336")
    C(92, "=MIN(D81,D6+D21+D35+D53+D72+D141)", "Rentas exentas y deducciones imputables limitadas (28+41+53+69+86+139)",
      "Todos los beneficios aceptados: limitados + 1% facturas + 72 UVT por dependientes.", "Instructivo c.92")
    C(93, "=MAX(0,D81-D82)", "Renta líquida ordinaria cédula general (91-92)", "", "")
    C(94, f'={R(BF(27))}', "Compensaciones por pérdidas 2018 y anteriores", "Requiere contador.", "Art. 147")
    C(95, f'={R(BF(28))}', "Compensaciones por exceso de renta presuntiva", "", "Art. 189")
    C(96, f'={R(BF(29))}', "Rentas gravables (activos omitidos/pasivos inexistentes)", "", "Art. 239-1")
    C(97, "=MAX(0,D83-D84-D85+D86)", "Renta líquida gravable cédula general (93-94-95+96)", "", "")
    C(98, f'={R(f"MAX(0,D9*{PAR}!$C$25+{ESP(64)})")}', "Renta presuntiva", "Tarifa 0% desde 2021 (parametrizada).", "Arts. 188-189")
    C(99, f'={R(f"{ING}!$F$33")}', "Ingresos brutos por pensiones", "", "Art. 337")
    C(100, f'={R(BF(30))}', "Ingresos no constitutivos de renta", "", "")
    C(101, "=MAX(0,D91-D92)", "Renta líquida (99-100)", "", "")
    C(102, f'={R(f"MIN(D93,{BF(31)}*1000*{PAR}!$C$5)")}', "Rentas exentas de pensiones", "Hasta 1.000 UVT por mes.", "Art. 206 num. 5")
    C(103, "=MAX(0,D93-D94)", "Renta líquida gravable de pensiones (101-102)", "", "")
    C(104, f'={R(BF(32))}', "Dividendos año 2016 y anteriores", "Régimen de transición.", "Arts. 48-49")
    C(105, f'={R(f"MIN(D98,{ESP(61)})")}', "Ingresos no constitutivos de renta", "", "")
    C(106, "=MAX(0,D98-D99)", "Renta líquida ordinaria año 2016 y anteriores (104-105)", "", "")
    C(107, f'={R(f"{ING}!$F$34")}', "1.ª subcédula dividendos 2017+", "Va a la tabla del 241 con descuento 254-1.", "Art. 242")
    C(108, f'={R(f"{ING}!$F$35")}', "2.ª subcédula dividendos 2017+", "Primero 35% y el neto a la tabla.", "Arts. 240 y 242")
    C(109, f'={R(ESP(62))}', "Dividendos recibidos del exterior", "Requiere contador.", "")
    C(110, f'={R(f"MIN(D103,{ESP(63)})")}', "Rentas exentas de la casilla 109", "", "")
    C(111, f'={R("MAX(D87,D88)+D95+D101+D102-D116")}', "RENTA LÍQUIDA GRAVABLE (97 o 98 + 103 + 107 + 108 − 118)",
      "La base del impuesto de la tabla del art. 241.", "Art. 331")
    C(112, f'={R(f"{ING}!$F$36+{ING}!$F$37")}', "Ingresos por ganancias ocasionales", "", "Arts. 299-306")
    C(113, f'={R(BF(33))}', "Costos por ganancias ocasionales", "", "")
    C(114, f'={R(f"MIN(MAX(0,D108-D109),{BF(34)})")}', "Ganancias ocasionales no gravadas y exentas", "", "Arts. 307 y 311-1")
    C(115, "=MAX(0,D108-D109-D110)", "Ganancias ocasionales gravables (112-113-114)", "", "")
    C(116, f'=IF(D87>=D88,{tarifa_fx("D105", T241, o)},0)', "Impuesto — tabla art. 241 (cédula general y pensiones)",
      "Base en UVT truncada a 2 decimales; épsilon 1E-9 contra errores de coma flotante en los bordes.", "Art. 241")
    C(117, f'=IF(D88>D87,{tarifa_fx("D105", T241, o)},0)', "Impuesto sobre renta presuntiva", "", "")
    C(118, f'={R(f"D102*{PAR}!$C$24")}', "Impuesto dividendos 2.ª subcédula (35%)", "", "Art. 240")
    C(119, f'={tarifa_fx("D100", T2016, o)}', "Impuesto dividendos 2016 y anteriores", "Tabla histórica.", "")
    C(120, f'={R("MAX(0,D103-D104)*0.35")}', "Impuesto dividendos del exterior", "", "")
    C(121, "=SUM(D114:D118)", "TOTAL impuesto sobre las rentas líquidas gravables (116 a 120)", "", "")
    C(122, f'={R(BF(36))}', "Descuento — impuestos pagados en el exterior", "", "Art. 254")
    C(123, f'={R(BF(37))}', "Descuento — donaciones (25% de lo donado)", "", "Arts. 257-258")
    C(124, f'={R(f"IF(TRUNC(MAX(0,D101+D102-D116)/{PAR}!$C$5,2)<=1090,0,TRUNC((TRUNC(MAX(0,D101+D102-D116)/{PAR}!$C$5,2)-1090)*0.19+0.000000001,2)*{PAR}!$C$5)+{BF(38)}")}',
      "Descuento — dividendos (art. 254-1) y otros", "", "Art. 254-1")
    C(125, "=MIN(SUM(D120:D122),D119)", "Total descuentos tributarios (122+123+124)", "Nunca superan el impuesto.", "Art. 259")
    C(126, "=MAX(0,D119-D123)", "IMPUESTO NETO DE RENTA (121-125)", "", "")
    lotg = f'MIN(D111,{R(f"{ING}!$F$37")})'
    C(127, f'={R(f"(D111-{lotg})*{PAR}!$C$22+{lotg}*{PAR}!$C$23")}', "Impuesto de ganancias ocasionales",
      "15% general; 20% loterías.", "Arts. 314 y 317")
    C(128, f'={R(f"MIN(D125,{BF(35)})")}', "Descuento impuesto exterior sobre GO", "", "")
    C(129, "=MAX(0,D124+D125-D126)", "TOTAL IMPUESTO A CARGO (126+127-128)", "", "")
    C(130, f'={R(f"{PER}!$B$21")}', "Anticipo liquidado año gravable anterior", "Casilla 133 de su declaración 2024.", "Art. 807")
    C(131, f'={R(f"{PER}!$B$20")}', "Saldo a favor año anterior sin devolución", "Casilla 137 de 2024.", "")
    C(132, f'={R(f"{ING}!$F$38")}', "Retenciones año gravable a declarar", "", "Arts. 365 y ss.")
    C(133, "=MIN(D133,D134)", "Anticipo renta año siguiente (menor de las dos opciones)", "", "Art. 807")
    C(134, "=MAX(0,D127+D135-D130-D131-D132)", "Saldo a pagar por impuesto", "", "")
    C(135, f'={R(BF(39))}', "Sanciones", "Mínimo 10 UVT de 2026 si aplica.", "Arts. 639-641")
    C(136, "=MAX(0,D127+D135+D137-D130-D131-D132)", "TOTAL SALDO A PAGAR", "", "")
    C(137, "=MAX(0,D130+D131+D132-D127-D135-D137)", "TOTAL SALDO A FAVOR", "Pagó de más: devolución/compensación o arrastre a 2026.", "Arts. 850-855")
    C(138, f"={DEP}!$B$17", "Número de dependientes económicos", "", "")
    C(139, f'={R(f"MIN(MAX(0,D21+D22+D35+D39),{BF(22)})")}', "Adición por dependientes a la casilla 92 (72 UVT c/u)",
      "FUERA del límite del 40%; tope: rentas líquidas de trabajo (regla Q53 del AyudaRenta).", "Art. 336 num. 3 inc. 2")
    C(140, '=IF(D25=0,"",IF(D27/D25>0.6,"X",""))', "¿Superó tope indicativo de costos 60%? (X)", "", "Art. 336-1")
    C(141, f'={R(BF(40))}', "Aporte voluntario", "", "Art. 244-1")
    C(980, "=D138+D143", "PAGO TOTAL con este formulario", "", "")
    return s

SECCIONES = {5: "PATRIMONIO Y FACTURA ELECTRÓNICA (28-31)", 11: "CÉDULA GENERAL — RENTAS DE TRABAJO (32-42)",
             24: "TRABAJO SIN RELACIÓN LABORAL, CON COSTOS (43-57)", 41: "RENTAS DE CAPITAL (58-73)",
             59: "RENTAS NO LABORALES (74-90)", 78: "RESUMEN CÉDULA GENERAL (91-98)",
             90: "CÉDULA DE PENSIONES (99-103)", 97: "DIVIDENDOS Y PARTICIPACIONES (104-111)",
             107: "GANANCIAS OCASIONALES (112-115)", 113: "LIQUIDACIÓN PRIVADA (116-129)",
             129: "SALDOS, RETENCIONES Y ANTICIPO (130-141)", 144: "PAGO"}
AUX_SPEC = {
    79: ("AUX", "Base del límite (ingresos + ECE − INCRNGO)", "=MAX(0,D12+D25+D42+D60+D46+D65-D13-D26-D43-D62)",
         "Art. 336 num. 1-2: la base NO resta costos.", ""),
    80: ("AUX", "Límite combinado 40% / 1.340 UVT",
         f"=ROUND(MIN(0.4*D79,{PAR}!$C$7*{PAR}!$C$5),-3)", "Se reparte en cascada entre subcédulas.", "Art. 336 num. 3"),
    133: ("AUX", "Anticipo — opción impuesto actual",
          f'=MAX(0,ROUND(IF({PER}!$B$18=1,0.25,IF({PER}!$B$18=2,0.5,0.75))*D124-D132,-3))', "", "Art. 807"),
    134: ("AUX", "Anticipo — opción promedio dos años",
          f'=IF({PER}!$B$18=1,D133,MAX(0,ROUND(IF({PER}!$B$18=2,0.5,0.75)*(D124+{PER}!$B$19)/2-D132,-3)))', "", "Art. 807"),
}

def hoja_calculo(nombre, oficial, ttl, nota):
    ws = wb.create_sheet(nombre)
    widths(ws, {"A": 8, "B": 44, "C": 3, "D": 16, "E": 52, "F": 22})
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    titulo(ws, "ABDEF", ttl, nota)
    header_row(ws, 4, "ABDEF", ["Casilla", "Concepto", "Valor", "En palabras simples", "Norma"])
    s = spec(oficial)
    for row, txt in SECCIONES.items():
        sec_row(ws, row, txt, "ABDEF")
    for fila, (cas, concepto, formula, simple, norma) in s.items():
        put(ws, f"A{fila}", cas, S_cas)
        put(ws, f"B{fila}", concepto, S_txt, sz=8)
        fmt = "@" if cas == "140" else "#,##0" + ("" if oficial else ".00")
        tot = cas in ("29", "31", "41", "91", "92", "97", "111", "116", "121", "126", "129", "132", "136", "137", "139", "980")
        put(ws, f"D{fila}", formula, S_fx, fmt=fmt, tot=tot)
        put(ws, f"E{fila}", simple, S_note)
        put(ws, f"F{fila}", norma, S_note)
        ws.row_dimensions[fila].height = max(18, 11 * (1 + len(simple) // 65))
    for fila, (cas, concepto, formula, simple, norma) in AUX_SPEC.items():
        put(ws, f"A{fila}", cas, S_note)
        put(ws, f"B{fila}", concepto, S_txt, sz=8, it=True)
        put(ws, f"D{fila}", formula, S_fx, fmt="#,##0")
        put(ws, f"E{fila}", simple, S_note); put(ws, f"F{fila}", norma, S_note)
    return ws

hoja_calculo(SN["F210"], True,
             "PASO 6 · FORMULARIO 210 — CASILLAS OFICIALES, aproximadas a miles (art. 577 E.T.)",
             "Copie cada valor azul en la casilla del MISMO número en el MUISCA. Estructura de la Resolución DIAN 000044/2024. No edite fórmulas: las entradas viven en 03_PERFIL, 04_INGRESOS, 05_BENEFICIOS, 06_DEPENDIENTES y 07_PATRIMONIO.")
hoja_calculo(SN["MX"], False,
             "MOTOR EXACTO — la misma cadena de cálculo SIN redondeo a miles (trazabilidad al peso)",
             "Capa analítica: valores antes de la aproximación del art. 577. Aquí se ve el efecto de $1 en las entradas. No se transcribe al MUISCA.")

# ============================================================ 11_LIMITE_40
ws = wb.create_sheet(SN["LIM"])
widths(ws, {"A": 64, "B": 16, "C": 52})
ws.sheet_view.showGridLines = False
titulo(ws, "ABC", "EL LÍMITE DEL 40% / 1.340 UVT — la regla que más define el impuesto de la cédula general",
       "Todas las rentas exentas y deducciones de la cédula general compiten por un solo techo: 40% del ingreso neto, máximo 1.340 UVT ($66.731.000). Fuera del techo solo quedan: 72 UVT por dependiente (c.139), 1% de facturas (c.28) y exenciones especiales (E-21).")
nota_paso2 = CFG.get("NOTA_LIMITE", "Compare el PASO 3 con el PASO 2: si lo pedido supera el techo, hay RECORTE y conseguir más certificados de esos rubros NO baja el impuesto.")
lims = [
    ("PASO 1 — Base: ingresos de la cédula general (+ECE) menos INCRNGO", f"={F210}!D79", "Art. 336 num. 1-2. Los costos NO se restan de la base."),
    ("PASO 2 — Techo: el MENOR entre 40% de la base y 1.340 UVT", f"={F210}!D80", nota_paso2),
    ("PASO 3 — Lo que usted pide (AFC + 25% + intereses + salud + dep. 387)", f"={F210}!D17+{F210}!D20+{F210}!D31+{F210}!D34+{F210}!D49+{F210}!D52+{F210}!D68+{F210}!D71-{BEN}!$E$65", "Suma de exentas y deducciones sujetas al límite."),
    ("PASO 4 — Lo que la ley ACEPTA (casillas 41+53+69+86)", f"={F210}!D21+{F210}!D35+{F210}!D53+{F210}!D72", "Asignación en cascada: trabajo → honorarios → capital → no laborales."),
    ("RECORTE: beneficios pedidos que el límite eliminó", "=MAX(0,B7-B8)", "Si es > 0, conseguir más certificados de estos rubros NO baja su impuesto."),
    ("PASO 5 — Beneficios FUERA del límite (c.28 + c.139 + E-21)", f"={F210}!D6+{F210}!D141+{BEN}!$E$65", "Los únicos que siguen sumando cuando el techo se copa."),
    ("Uso del límite", "=IF(B6=0,0,B8/B6)", "Porcentaje del techo consumido."),
]
r = 5
for lab, fx, com in lims:
    put(ws, f"A{r}", lab, S_txt, b=lab.startswith(("RECORTE", "Uso")))
    put(ws, f"B{r}", fx, S_fx, fmt="0.0%" if "Uso" in lab else "#,##0", tot=("PASO 2" in lab or "RECORTE" in lab))
    put(ws, f"C{r}", com, S_note)
    ws.row_dimensions[r].height = 24; r += 1
put(ws, "A13", '=IF(B9>0,"LECTURA ESTRATÉGICA: el límite le recortó $"&TEXT(B9,"#,##0")&". Más AFC, más medicina o el GMF ya NO mueven su impuesto. Solo suman: dependientes (72 UVT), 1% de facturas, RAIS voluntario (INCRNGO) y las exenciones especiales.","Sus beneficios caben completos dentro del límite.")', S_warn)
ws.merge_cells("A13:C14"); ws.row_dimensions[13].height = 26

# ============================================================ 12_CONTROLES
ws = wb.create_sheet(SN["CTRL"])
widths(ws, {"A": 52, "B": 30, "C": 11, "D": 34, "E": 44})
ws.sheet_view.showGridLines = False
titulo(ws, "ABCDE", "PASO 7 · CONTROLES DE CALIDAD Y DECISIÓN DE PRESENTACIÓN",
       "Un resultado calculado NO es un resultado presentable. Cierre los controles críticos antes de transcribir al MUISCA.")
header_row(ws, 4, "ABCDE", ["Control", "Resultado", "Severidad", "Evidencia", "Acción requerida"])
ctr = [
    ("Residencia fiscal (aplica Formulario 210)", f'=IF({PER}!B16="Sí","OK","BLOQUEO")', "Crítica", "03_PERFIL B16", "Si no es residente, use el 110"),
    ("Identidad completa contra el RUT", f'=IF(AND({PER}!B11<>"",{PER}!B13<>"",{PER}!B14<>""),"OK","FALTA")', "Crítica", "AyudaRenta 2024", "Verifique nombre, DV y seccional en el RUT"),
    ("Obligado a declarar", f'=IF(OR({FD(29)}>{PAR}!$C$18*{C5},{ING}!F27>={PAR}!$C$19*{C5},{CFG["movs_tope"]}>{PAR}!$C$19*{C5}),"SÍ — OBLIGADO","No por topes")', "Informativa", f"Arts. 592-594-3; exógena (movimientos ${CFG['movs_tope']:,.0f})".replace(",", "."), "Debe declarar"),
    ("Firma de contador", f'=IF(OR({FD(29)}>{PAR}!$C$20*{C5},{ING}!F27>{PAR}!$C$20*{C5}),"SÍ","No")', "Informativa", "Art. 596", ""),
    ("Fecha prevista dentro del plazo", f'=IF({PER}!B25<={PER}!B10,"OK","FUERA DE PLAZO")', "Crítica", "03_PERFIL B25 vs B10", f"Presente a más tardar el {CFG['fecha_limite']}"),
    ("Coherencia patrimonial (arts. 236-239)", f'=IF({FD(31)}-{PER}!B24<={FD(91)}+{FD(92)}+{FD(33)}+{FD(59)}-{FD(129)},"JUSTIFICADA","NO JUSTIFICADA — revise")', "Crítica", "Capacidad de la renta vs Δ patrimonio líquido", "Si no justifica, hay ingreso omitido o patrimonio mal valorado"),
    ("136 y 137 excluyentes", f'=IF(AND({FD(136)}>0,{FD(137)}>0),"ERROR","OK")', "Crítica", "Automático", ""),
    ("Casos especiales fuera de alcance en cero", f'=IF(SUM({BEN}!E49,{BEN}!E56,{BEN}!E62)+{BEN}!F27+{BEN}!F28+{BEN}!F29=0,"OK","REQUIERE CONTADOR")', "Crítica", "ECE, pérdidas, activos omitidos, div. exterior", ""),
    ("Paridad Python (QA) vigente", f'=IF({QAs}!B6="","PENDIENTE EJECUTAR NOTEBOOK",{QAs}!B6)', "Crítica", "Hoja 13_QA", "Corra validar.py tras cualquier cambio"),
    ("Fronteras de tarifa correctas en este libro", f'=IF({QAs}!B7=0,"OK","FALLA EN "&{QAs}!B7&" CASOS")', "Crítica", "13_QA filas 12-34", "Si falla, no use el archivo"),
]
ctr += CFG.get("CONTROLES_EXTRA", [])
r = 5
for lab, fx, sev, evi, acc in ctr:
    put(ws, f"A{r}", lab, S_txt, sz=8)
    put(ws, f"B{r}", fx if fx.startswith("=") else fx.strip('"'), S_fx, fmt="@")
    put(ws, f"C{r}", sev, S_note); put(ws, f"D{r}", evi, S_note); put(ws, f"E{r}", acc, S_note)
    ws.row_dimensions[r].height = 24; r += 1
FIN = r
sec_row(ws, FIN + 1, "DICTAMEN AUTOMÁTICO", "ABCDE")
put(ws, f"A{FIN+2}", "Controles con BLOQUEO / FALTA / ERROR / FUERA DE PLAZO / INVESTIGAR", S_txt)
put(ws, f"B{FIN+2}", f'=SUMPRODUCT(--(ISNUMBER(SEARCH("BLOQUEO",B5:B{FIN-1}))))+SUMPRODUCT(--(ISNUMBER(SEARCH("FALTA",B5:B{FIN-1}))))+SUMPRODUCT(--(ISNUMBER(SEARCH("ERROR",B5:B{FIN-1}))))+SUMPRODUCT(--(ISNUMBER(SEARCH("FUERA",B5:B{FIN-1}))))+SUMPRODUCT(--(ISNUMBER(SEARCH("INVESTIGAR",B5:B{FIN-1}))))', S_fx, fmt="0")
put(ws, f"A{FIN+3}", "Controles con REVISAR / ESTIMADO / PENDIENTE / VERIFICAR / DECISIÓN / DISMINUCIÓN", S_txt)
put(ws, f"B{FIN+3}", f'=SUMPRODUCT(--(ISNUMBER(SEARCH("REVISAR",B5:B{FIN-1}))))+SUMPRODUCT(--(ISNUMBER(SEARCH("ESTIMADO",B5:B{FIN-1}))))+SUMPRODUCT(--(ISNUMBER(SEARCH("PENDIENTE",B5:B{FIN-1}))))+SUMPRODUCT(--(ISNUMBER(SEARCH("VERIFICAR",B5:B{FIN-1}))))+SUMPRODUCT(--(ISNUMBER(SEARCH("DECISIÓN",B5:B{FIN-1}))))+SUMPRODUCT(--(ISNUMBER(SEARCH("DISMINUCIÓN",B5:B{FIN-1}))))', S_fx, fmt="0")
put(ws, f"A{FIN+4}", "ESTADO", S_txt, b=True)
put(ws, f"B{FIN+4}", f'=IF(B{FIN+2}>0,"NO LISTO PARA PRESENTAR",IF(B{FIN+3}>0,"CALCULADO — CIERRE PENDIENTES","LISTO PARA REVISIÓN FINAL"))', S_fx, fmt="@", tot=True)
DICT_ROWS = (FIN + 2, FIN + 3, FIN + 4)
put(ws, f"A{FIN+6}", "Calculadora de sanción por extemporaneidad: meses (o fracción) de retraso →", S_txt)
put(ws, f"B{FIN+6}", 0, S_in, fmt="0")
put(ws, f"A{FIN+7}", "Sanción estimada (5%/mes del impuesto a cargo, máx. 100%, mínimo 10 UVT 2026)", S_txt)
put(ws, f"B{FIN+7}", f'=IF(B{FIN+6}=0,0,MAX(MIN(0.05*B{FIN+6},1)*{FD(129)},ROUND({PAR}!$C$17*{PAR}!$C$6,-3)))', S_fx)

# ============================================================ 13_QA
ws = wb.create_sheet(SN["QA"])
widths(ws, {"A": 40, "B": 20, "C": 18, "D": 18, "E": 12, "F": 40})
ws.sheet_view.showGridLines = False
titulo(ws, "ABCDEF", "QA — PRUEBAS EMBEBIDAS Y RESULTADO DEL AUDITOR EXTERNO")
put(ws, "A3", "1. Resultado de la validación externa (péguelo tras cada corrida)", S_sub); ws.merge_cells("A3:F3")
put(ws, "A4", "Fecha y estado de la última auditoría Python", S_txt)
put(ws, "B4", "PENDIENTE — ejecute scripts/validar.py", S_in, fmt="@")
put(ws, "A5", "Regla: si cambió CUALQUIER entrada o fórmula después de esa fecha, vuelva a correr la validación antes de presentar.", S_note)
ws.merge_cells("A5:F5")
put(ws, "A6", "Estado usado por 12_CONTROLES", S_txt); put(ws, "B6", "=B4", S_fx, fmt="@")
put(ws, "A7", "Fronteras con FALLA (de la sección 2)", S_txt)
put(ws, "B7", "=COUNTIF(E12:E34,\"FALLA\")", S_fx, fmt="0")
put(ws, "A9", "2. Fronteras de la tarifa art. 241 — el libro se prueba a sí mismo", S_sub); ws.merge_cells("A9:F9")
put(ws, "A10", "Cada fila aplica la MISMA fórmula de tarifa del libro a una base de prueba y la compara con el valor esperado (calculado de forma independiente en Python con Decimal).", S_note)
ws.merge_cells("A10:F10")
header_row(ws, 11, "ABCDE", ["Caso", "Base (COP)", "Impuesto del libro", "Esperado (Python)", "Veredicto"])
uvt = D(49799)
bordes = []
for fUVT in [D("1090"), D("1700"), D("4100"), D("8670"), D("18970"), D("31000")]:
    for db in [D("-0.01"), D("0"), D("0.01")]:
        bordes.append(fUVT + db)
bordes += [D("500"), D("2500.37"), D("12000.5"), D("50000"), D("3603.72")]
r = 12
for buvt in bordes:
    base = (uvt * buvt).quantize(D("0.01"))
    esperado = _tarifa(base, P.tabla241, uvt, True)
    put(ws, f"A{r}", f"base {buvt} UVT", S_note)
    c = ws[f"B{r}"]; c.value = float(base); c.number_format = "#,##0.00"; c.font = Font(name=F, size=8)
    put(ws, f"C{r}", "=" + tarifa_fx(f"B{r}", T241, True), S_fx)
    c = ws[f"D{r}"]; c.value = int(esperado); c.number_format = "#,##0"; c.font = Font(name=F, size=8)
    put(ws, f"E{r}", f'=IF(C{r}=D{r},"OK","FALLA")', S_fx, fmt="@")
    r += 1
put(ws, f"A{r+1}", "3. Nota metodológica", S_sub); ws.merge_cells(f"A{r+1}:F{r+1}")
put(ws, f"A{r+2}", "Convención: base en UVT truncada a 2 decimales; producto marginal truncado a 2 decimales con épsilon 1E-9 (blindaje contra coma flotante binaria de Excel); resultado al múltiplo de mil (art. 577). Es la convención del AyudaRenta oficial.", S_note)
ws.merge_cells(f"A{r+2}:F{r+3}")

# ============================================================ 14_SOPORTES
ws = wb.create_sheet(SN["SOP"])
widths(ws, {"A": 4, "B": 44, "C": 14, "D": 10, "E": 34, "F": 30, "G": 36})
ws.sheet_view.showGridLines = False
titulo(ws, "ABCDEFG", "SOPORTES — CARPETA DOCUMENTAL (consérvelos mínimo 5 años)")
header_row(ws, 3, "ABCDEFG", ["N.º", "Documento", "Estado (edite)", "Criticidad", "Para qué se usa", "Acción", "Observación"])
sops = CFG.get("SOPORTES") or [
    ("RUT actualizado", "FALTA", "Crítica", "Identidad y responsabilidades", "Descargar de la DIAN", ""),
    ("Declaración del año anterior (PDF presentado)", "FALTA", "Alta", "Saldos de arrastre e histórico", "Descargar con constancia", ""),
    ("Información exógena del año gravable", "FALTA", "Crítica", "Conciliación por tercero", "Descargar del MUISCA", ""),
    ("Certificado 220 y certificados bancarios de rendimientos/retenciones", "FALTA", "Crítica", "Ingresos y casilla 132", "Pedir a cada entidad", ""),
    ("Certificados de saldos, créditos, AFC, medicina, dependientes", "FALTA", "Crítica", "Patrimonio y beneficios", "Pedir a cada entidad", ""),
]
for i, (doc, est, cri, uso, acc, obs) in enumerate(sops):
    r = 4 + i
    put(ws, f"A{r}", i + 1, S_note); put(ws, f"B{r}", doc, S_txt, sz=8)
    put(ws, f"C{r}", est, S_in, fmt="@")
    put(ws, f"D{r}", cri, S_note); put(ws, f"E{r}", uso, S_note)
    put(ws, f"F{r}", acc, S_note); put(ws, f"G{r}", obs, S_note)
    ws.row_dimensions[r].height = 20
rr = 5 + len(sops)
put(ws, f"B{rr}", 'Marque "OK" en la columna C cuando tenga el documento. Regla de oro: sin soporte, el valor no se declara (art. 786 E.T. — la carga de la prueba es suya).', S_alert)
ws.merge_cells(f"B{rr}:G{rr}"); ws.row_dimensions[rr].height = 24

# ============================================================ 15_EXOGENA
ws = wb.create_sheet(SN["EXO"])
widths(ws, {"A": 12, "B": 40, "C": 46, "D": 16, "E": 40})
ws.sheet_view.showGridLines = False
titulo(ws, "ABCDE", "INFORMACIÓN EXÓGENA 2025 — COPIA DE REFERENCIA (no editar; no es certificado)")
header_row(ws, 3, "ABCDE", ["NIT reporta", "Quién reporta", "Concepto", "Valor", "Uso sugerido DIAN"])
exo = CFG.get("EXOGENA", [])
topes = CFG.get("EXO_TOPES", "Registre aquí los topes que muestra el reporte de exógena de la DIAN (ingresos, patrimonio, consumos, movimientos, compras).")
for i, (nit, quien, conc, val, uso) in enumerate(exo):
    r = 4 + i
    put(ws, f"A{r}", nit, S_note); put(ws, f"B{r}", quien, S_txt, sz=8)
    put(ws, f"C{r}", conc, S_txt, sz=8)
    c = ws[f"D{r}"]; c.value = val; c.number_format = "#,##0"; c.font = Font(name=F, size=8); c.alignment = Alignment(horizontal="right")
    put(ws, f"E{r}", uso, S_note)
rr = 5 + len(exo)
put(ws, f"A{rr}", topes + " La exógena es orientativa: la declaración debe reflejar su realidad económica (advertencia de la propia DIAN).", S_note)
ws.merge_cells(f"A{rr}:E{rr+1}")

# ============================================================ 17_DECISIONES
ws = wb.create_sheet(SN["DEC"])
widths(ws, {"A": 24, "B": 28, "C": 34, "D": 38, "E": 34})
ws.sheet_view.showGridLines = False
titulo(ws, "ABCDE", "MATRIZ DE DECISIONES — QUÉ CAMBIÓ EN v4 Y POR QUÉ (jerarquía de evidencia)",
       "Cada decisión se tomó por evidencia documental: instructivo oficial > AyudaRenta oficial > declaración presentada > exógena > estimación. Nada se eligió por conveniencia numérica. Las decisiones v1/v2/v3 previas se conservan en el archivo v3.")
header_row(ws, 4, "ABCDE", ["Tema", "Versión anterior (v3)", "Evidencia nueva", "DECISIÓN v4", "Norma / soporte"])
dec = CFG.get("DECISIONES", [])
for i, row in enumerate(dec):
    r = 5 + i
    tema, v3v, ev, v4v, norma = row
    put(ws, f"A{r}", tema, S_txt, sz=8, b=True); put(ws, f"B{r}", v3v, S_note)
    put(ws, f"C{r}", ev, S_note); put(ws, f"D{r}", v4v, S_txt, sz=8)
    put(ws, f"E{r}", norma, S_note)
    ws.row_dimensions[r].height = 40
    if i % 2:
        for cc in "ABCDE": ws[f"{cc}{r}"].fill = PatternFill("solid", fgColor="F2F2F2")

# ============================================================ 18_GLOSARIO
ws = wb.create_sheet(SN["GLO"])
widths(ws, {"A": 26, "B": 58, "C": 56})
ws.sheet_view.showGridLines = False
titulo(ws, "ABC", "GLOSARIO — CADA TÉRMINO EN PALABRAS SIMPLES Y EN PRECISIÓN TÉCNICA")
put(ws, "A3", "Término", S_sub); put(ws, "B3", "En palabras simples", S_sub); put(ws, "C3", "Precisión técnica y norma", S_sub)
glos = [
    ("Año gravable / UVT", "El año de los ingresos que declara (2025). La UVT es la 'unidad de medida' de todos los topes.",
     "UVT 2025 = $49.799 (Res. 000193/2024); UVT 2026 = $52.374 (Res. 000238/2025). Topes del año gravable con la primera; sanciones pagadas en 2026 con la segunda."),
    ("Formulario 210", "El formato oficial que llenan las personas naturales residentes.", "Res. DIAN 000044/2024; se diligencia y firma en el MUISCA (Res. 000227/2025: 100% virtual)."),
    ("MUISCA", "La página de la DIAN donde se llena y firma la declaración.", "Sistema de la DIAN; el recibo de pago es el Formulario 490."),
    ("RUT", "Su 'cédula tributaria': ahí están su NIT, actividad y responsabilidades.", "Registro Único Tributario (art. 555-2 E.T.)."),
    ("Exógena", "Lo que los bancos, empleadores y demás le contaron a la DIAN sobre usted.", "Información de terceros (arts. 631 y ss.). Es orientativa: prevalece su realidad económica con soportes."),
    ("Casilla", "Cada número del formulario (c.29, c.32…). Este archivo calcula todas.", "Numeración oficial del instructivo del 210."),
    ("Cédula general", "La 'canasta' donde se suman trabajo, capital y no laborales.", "Arts. 330-336 E.T.; cada tipo de renta se depura por subcédula."),
    ("INCRNGO", "Ingresos que la ley saca de la base ANTES de calcular (no pagan impuesto).", "Ingreso No Constitutivo de Renta ni Ganancia Ocasional: aportes obligatorios a salud/pensión (arts. 55-56), componente inflacionario (arts. 38-41)."),
    ("Renta exenta / deducción / descuento", "Tres formas de pagar menos: las dos primeras restan de la BASE; el descuento resta del IMPUESTO (vale más).",
     "Exentas y deducciones: casillas 35-39; descuentos: 122-124 (límite art. 259)."),
    ("Límite 40% / 1.340 UVT", "El techo que junta casi todos los beneficios: por más papeles que consiga, no pasa de ahí.",
     "Art. 336 num. 3. Por fuera: 72 UVT por dependiente, 1% de facturas, RAIS voluntario (INCRNGO) y exenciones especiales."),
    ("Renta exenta del 25%", "La cuarta parte del salario depurado no paga impuesto.", "Art. 206 num. 10; tope 790 UVT; NO aplica sobre cesantías (instructivo c.36)."),
    ("Cesantías", "Se declaran cuando se pagan o consignan; tienen SU PROPIA exención.", "Art. 206 num. 4: tabla según salario promedio de 6 meses (en 16_PARAMETROS)."),
    ("AFC / FVP / AVC", "Cuentas de ahorro con beneficio si el dinero va a vivienda o se queda 10 años.", "Arts. 126-1/126-4; tope conjunto 30% del ingreso y 3.800 UVT."),
    ("Componente inflacionario", "La parte de los intereses que solo compensa la inflación: no paga impuesto.",
     "55,43% de los rendimientos brutos AG2025 no gravado; espejo: 28,35% de los costos financieros no deducible (Dcto. 898/2026)."),
    ("Dependiente económico", "Familiar que vive de usted; da DOS rebajas si usted tiene salario.",
     "Art. 387 (32 UVT/mes) + art. 336-3 (72 UVT/año, fuera del límite); ambas exigen renta laboral (DUR 1.2.1.20.3)."),
    ("Patrimonio bruto / líquido", "Todo lo que tiene a 31/12 / lo mismo menos las deudas.", "Arts. 261-282; inmuebles al MAYOR entre costo fiscal y avalúo (art. 277); deudas con soporte de fecha cierta (arts. 283, 767, 770)."),
    ("Avalúo catastral vs costo fiscal", "El valor del predial vs lo que le costó (con ajustes). Se declara el MAYOR.", "Art. 277 E.T. (no obligados a contabilidad)."),
    ("Comparación patrimonial", "Si su patrimonio creció más de lo que sus rentas explican, la DIAN grava la diferencia.",
     "Arts. 236-239. Una DISMINUCIÓN no genera renta, pero conviene poder explicarla."),
    ("Retención en la fuente", "Impuesto que le descontaron por adelantado durante el año.", "Arts. 365 y ss.; va en la casilla 132 y se resta del impuesto."),
    ("Anticipo", "Abono obligatorio al impuesto del AÑO SIGUIENTE.", "Art. 807: 25%/50%/75% según años declarando, con opción de promedio; se resta la retención."),
    ("Saldo a favor", "Pagó de más: puede pedir devolución, compensar o arrastrarlo.", "Arts. 850-855 E.T.; SIE de devoluciones de la DIAN."),
    ("Estimado (celda ámbar)", "Valor provisional puesto con la MEJOR evidencia disponible; usted debe reemplazarlo con el certificado.",
     "Convención de este archivo. El control de 12_CONTROLES no da 'LISTO' mientras existan ámbar sin resolver."),
    ("Corrección de declaraciones", "Arreglar una declaración ya presentada.", "Art. 588 (dentro de los 3 años; si no cambia el impuesto no hay sanción) y art. 589 (si disminuye el saldo a pagar)."),
    ("Firmeza", "El plazo tras el cual la DIAN ya no puede revisar la declaración.", "Regla general 3 años (art. 714); beneficio de auditoría 6/12 meses (art. 689-3) si el impuesto crece 35%/25%."),
    ("Beneficio de auditoría", "Su declaración queda 'en firme' más rápido si paga suficientemente más que el año pasado.",
     "Art. 689-3: firmeza en 6 meses (crecimiento del impuesto neto ≥35%) o 12 meses (≥25%); exige impuesto 2024 ≥ 71 UVT."),
    ("Sanciones", "Multas por presentar tarde, con errores o corregir.", "Extemporaneidad 5%/mes (arts. 641-642), mínima 10 UVT 2026 = $524.000 (art. 639); inexactitud 100% (arts. 647-648); corrección 10-20% (art. 644)."),
    ("Fuentes oficiales", "Dónde verificar todo lo anterior.",
     "E.T.: secretariasenado.gov.co/senado/basedoc/estatuto_tributario.html · Micrositio DIAN renta AG2025 · Dcto. 898/2026 · Res. 000044/2024 y 000120/2024 · Dcto. 2229/2023 (plazos)."),
]
for i, (t, simple, tec) in enumerate(glos):
    r = 4 + i
    put(ws, f"A{r}", t, S_txt, sz=8, b=True); put(ws, f"B{r}", simple, S_txt, sz=8)
    put(ws, f"C{r}", tec, S_note)
    ws.row_dimensions[r].height = max(24, 12 * (1 + max(len(simple) // 70, len(tec) // 68)))
    if i % 2:
        for cc in "ABC": ws[f"{cc}{r}"].fill = PatternFill("solid", fgColor="F2F2F2")

# ============================================================ 02_METODOLOGIA
ws = wb.create_sheet(SN["MET"])
widths(ws, {"A": 5, "B": 112})
ws.sheet_view.showGridLines = False
titulo(ws, "AB", "METODOLOGÍA — CÓMO FUNCIONA ESTE ARCHIVO Y CÓMO CONFIAR EN ÉL")
met = [
    ("SEC", "1. QUÉ ES ESTO"),
    ("TXT", "Un solo archivo que: (a) recoge sus datos en 5 hojas de entrada, (b) calcula las 114 casillas del Formulario 210 con las reglas del Estatuto Tributario, (c) se vigila a sí mismo con controles y pruebas embebidas, y (d) queda auditado por un programa Python independiente que rehace TODO el cálculo desde cero y exige coincidencia del 100%. Usted solo edita celdas VERDES y reemplaza las ÁMBAR."),
    ("SEC", "2. EL FLUJO EN 7 PASOS (el orden de las hojas ES el orden de trabajo)"),
    ("TXT", "PASO 1 · 03_PERFIL: quién declara y los saldos del año pasado.  PASO 2 · 04_INGRESOS: cada peso que recibió, una fila por tercero.  PASO 3 · 05_BENEFICIOS: lo que la ley le permite restar (con su regla y su tope automático).  PASO 4 · 06_DEPENDIENTES: quiénes dependen de usted.  PASO 5 · 07_PATRIMONIO: lo que tiene y lo que debe a 31/12, comparado con 2024.  PASO 6 · 09_FORMULARIO_210: el resultado casilla por casilla (no se toca: se lee y se transcribe).  PASO 7 · 12_CONTROLES: el dictamen — solo presente cuando esté en verde."),
    ("SEC", "3. CONVENCIONES DE COLOR (las mismas en todo el archivo)"),
    ("IN",  "VERDE — celda de ENTRADA: usted la edita."),
    ("EST", "ÁMBAR — valor ESTIMADO con la mejor evidencia disponible: DEBE reemplazarlo con el certificado del año."),
    ("FX",  "AZUL NEGRILLA — FÓRMULA: no la toque; se recalcula sola."),
    ("WRN", "NARANJA — advertencia que pide una acción suya."),
    ("ALR", "ROJO — bloqueo: no presente sin resolverlo."),
    ("SEC", "4. CÓMO CALCULA (dos motores en paralelo)"),
    ("TXT", "09_FORMULARIO_210 aplica el redondeo LEGAL al múltiplo de mil (art. 577 E.T.) en cada casilla — es lo que se transcribe al MUISCA. 10_MOTOR_EXACTO ejecuta la MISMA cadena sin redondear: ahí se ve el efecto de cada peso. La tarifa del art. 241 usa la convención oficial del AyudaRenta de la DIAN: base en UVT truncada a 2 decimales, épsilon 1E-9 contra errores de coma flotante, resultado a miles."),
    ("SEC", "5. CÓMO SE VERIFICA (tres candados independientes)"),
    ("TXT", "CANDADO 1 — Controles vivos (12_CONTROLES): residencia, plazos, coherencia patrimonial (arts. 236-239), cruces contra exógena, exclusión mutua de saldos. Producen un DICTAMEN automático. CANDADO 2 — Pruebas embebidas (13_QA): 23 casos de frontera de la tarifa del art. 241 calculados de forma independiente en Python quedan DENTRO del archivo comparándose en vivo contra las fórmulas del libro. CANDADO 3 — Auditor externo: un notebook Python reconstruye el caso leyendo SOLO las celdas de entrada, recalcula las 114 casillas con aritmética Decimal exacta, exige paridad 100% en ~230 comparaciones, corre 15 escenarios de estrés, una prueba de mutación (corrompe el archivo a propósito y verifica que la auditoría LO DETECTE) y una prueba de sensibilidad al peso."),
    ("SEC", "6. JERARQUÍA DE EVIDENCIA (qué manda cuando las fuentes pelean)"),
    ("TXT", "1.º Instructivo oficial del Formulario 210 (Res. 000044/2024) · 2.º AyudaRenta oficial DIAN · 3.º Certificados de terceros · 4.º Declaraciones presentadas de años anteriores · 5.º Exógena · 6.º Estimaciones (SIEMPRE en ámbar). Regla de estados: 'Confirmado' = dos fuentes coinciden o hay certificado; 'Provisional' = una fuente confiable sin certificado; 'ESTIMADO' = inferencia documentada que usted debe reemplazar; 'Falta/Pendiente' = no se declara sin soporte (art. 786: la carga de la prueba es suya)."),
    ("SEC", "7. QUÉ HACER SI CAMBIA UN VALOR"),
    ("TXT", "Edite la celda verde/ámbar → el libro recalcula solo → revise 12_CONTROLES (el dictamen cambia) → vuelva a correr el notebook auditor → registre fecha y resultado en 13_QA → recién entonces transcriba al MUISCA. Si un control queda en rojo, NO presente."),
    ("SEC", "8. MARCO NORMATIVO ESENCIAL"),
    ("TXT", "Depuración cedular: arts. 26, 330-343 · Rentas de trabajo: arts. 103, 206 · INCRNGO: arts. 55-56 · Capital: arts. 335, 38-41 (componente inflacionario, Dcto. 898/2026) · Beneficios: arts. 119, 126-1, 126-4, 387 · Límite conjunto: art. 336 · Tarifas: arts. 241, 314, 317 · Patrimonio: arts. 261-283 · Comparación patrimonial: arts. 236-239 · Anticipo: art. 807 · Redondeo: art. 577 · Sanciones: arts. 639-648 · Firmeza: arts. 714, 689-3 · Plazos: Dcto. 2229/2023."),
    ("SEC", "9. LO QUE ESTE ARCHIVO NO ES"),
    ("ALR", "No es un dictamen profesional ni asesoría legal: es una hoja de trabajo rigurosa. La coincidencia matemática Excel↔Python certifica el CÁLCULO, no los hechos económicos ni los soportes. La responsabilidad de lo declarado es del declarante (arts. 571, 596, 647-648 E.T.)."),
]
r = 3
for kind, txt in met:
    if kind == "SEC":
        sec_row(ws, r, txt, "AB"); ws.row_dimensions[r].height = 18
    else:
        st = {"TXT": S_txt, "IN": S_in, "EST": S_est, "WRN": S_warn, "ALR": S_alert}.get(kind)
        if kind == "FX":
            c = put(ws, f"B{r}", txt); c.font = Font(name=F, size=9, bold=True, color=BLU)
        elif kind in ("IN", "EST"):
            put(ws, f"B{r}", txt, st, fmt="@")
        else:
            put(ws, f"B{r}", txt, st)
        ws.merge_cells(f"B{r}:B{r}")
        est_h = max(16, 12 * (1 + len(txt) // 118))
        ws.row_dimensions[r].height = est_h
    r += 1
ws.column_dimensions["B"].width = 118

# ============================================================ 01_INICIO
ws = wb.create_sheet(SN["INICIO"])
widths(ws, {"A": 26, "B": 17, "C": 17, "D": 17, "E": 17, "F": 17, "G": 17, "H": 34})
ws.sheet_view.showGridLines = False
put(ws, "A1", f"DECLARACIÓN DE RENTA · AÑO GRAVABLE 2025 · FORMULARIO 210", S_title)
for cc in "BCDEFGH": put(ws, f"{cc}1", "", S_title)
ws.row_dimensions[1].height = 26
ver = CFG.get("VERSION_TXT", "Versión 1")
put(ws, "A2", f"{CFG['nombre']} · NIT {CFG['nit']}-{CFG['dv']} · {ver}", S_sub)
for cc in "BCDEFGH": put(ws, f"{cc}2", "", S_sub)
put(ws, "A3", "Fecha límite", S_txt, b=True)
put(ws, "B3", f"={PER}!B10", S_fx, fmt="dd/mm/yyyy", tot=True)
put(ws, "C3", "Días de margen:", S_txt, b=True)
put(ws, "D3", f"={PER}!B10-TODAY()", S_fx, fmt="0", tot=True)
put(ws, "E3", "DICTAMEN:", S_txt, b=True)
put(ws, "F3", f"={CTRL}!B{DICT_ROWS[2]}", S_fx, fmt="@", tot=True); ws.merge_cells("F3:H3")
kpis1 = [("Ingresos brutos totales", f"={ING}!F27"), ("Renta líquida gravable (c.111)", FD(111).replace("=", "=") if False else f"={F210}!D{fila_casilla(111)}"),
         ("IMPUESTO A CARGO (c.129)", f"={F210}!D{fila_casilla(129)}"), ("Retenciones + saldos previos", f"={F210}!D{fila_casilla(130)}+{F210}!D{fila_casilla(131)}+{F210}!D{fila_casilla(132)}")]
kpis2 = [("Patrimonio bruto (c.29)", f"={F210}!D{fila_casilla(29)}"), ("Patrimonio líquido (c.31)", f"={F210}!D{fila_casilla(31)}"),
         ("Variación patrimonio líquido vs 2024", f"={PAT}!E35"), ("Anticipo 2026 (c.133)", f"={F210}!D{fila_casilla(133)}")]
sec_row(ws, 5, "SU DECLARACIÓN EN OCHO CIFRAS (vivas: cambian si usted edita)", "ABCDEFGH")
cols = [("A", "B"), ("C", "D"), ("E", "F"), ("G", "H")]
for j, (lab, fx) in enumerate(kpis1):
    c1, c2 = cols[j]
    put(ws, f"{c1}6", lab, S_kpi_l); ws.merge_cells(f"{c1}6:{c2}6")
    put(ws, f"{c1}7", fx, S_kpi_v); ws.merge_cells(f"{c1}7:{c2}7")
for j, (lab, fx) in enumerate(kpis2):
    c1, c2 = cols[j]
    put(ws, f"{c1}9", lab, S_kpi_l); ws.merge_cells(f"{c1}9:{c2}9")
    put(ws, f"{c1}10", fx, S_kpi_v); ws.merge_cells(f"{c1}10:{c2}10")
    if "Variación" in lab: ws[f"{c1}10"].number_format = '"$" #,##0;[Red]"-$" #,##0'
ws.row_dimensions[7].height = 26; ws.row_dimensions[10].height = 26
fin_txt = CFG.get("FIN_TXT", "puede pedir devolución/compensación (arts. 850-855) o arrastrarlo al año siguiente (casilla 131).")
put(ws, "A12", f'=IF({F210}!D{fila_casilla(137)}>0,"RESULTADO: SALDO A SU FAVOR de $"&TEXT({F210}!D{fila_casilla(137)},"#,##0")&" — {fin_txt}","RESULTADO: SALDO A PAGAR de $"&TEXT({F210}!D{fila_casilla(136)},"#,##0")&" — pague con el Formulario 490 antes de la fecha límite.")', S_warn)
ws.merge_cells("A12:H13"); ws.row_dimensions[12].height = 22
pend = CFG.get("PENDIENTES_TXT", "LO QUE FALTA: revise 12_CONTROLES — allí está la lista viva de pendientes, su severidad y el dictamen automático. No presente con bloqueos ni con celdas ÁMBAR sin resolver.")
put(ws, "A15", pend, S_alert)
ws.merge_cells("A15:H17"); ws.row_dimensions[15].height = 20
sec_row(ws, 19, "RUTA DE TRABAJO — 7 PASOS (haga clic para ir a cada hoja)", "ABCDEFGH")
pasos = [
    ("03_PERFIL", "PASO 1 — Verifique identidad, años declarando y saldos 2024."),
    ("04_INGRESOS", "PASO 2 — Revise cada ingreso fila por fila; reemplace todo ámbar con su certificado."),
    ("05_BENEFICIOS", "PASO 3 — Active solo beneficios con soporte; reemplace los ámbar."),
    ("06_DEPENDIENTES", "PASO 4 — Dependientes con soporte (el archivo valida la regla laboral)."),
    ("07_PATRIMONIO", "PASO 5 — Confirme activos y deudas contra extractos; mire la columna de variación."),
    ("09_FORMULARIO_210", "PASO 6 — Lea el resultado casilla por casilla (10_MOTOR_EXACTO para trazar al peso)."),
    ("12_CONTROLES", "PASO 7 — Cierre pendientes hasta que el dictamen quede LISTO; luego MUISCA."),
]
for i, (hoja, texto) in enumerate(pasos):
    r = 20 + i
    put(ws, f"A{r}", f'=HYPERLINK("#\'{hoja}\'!A1","{hoja}")', S_fx, fmt="@")
    put(ws, f"B{r}", texto, S_txt); ws.merge_cells(f"B{r}:H{r}")
    ws.row_dimensions[r].height = 17
sec_row(ws, 28, "MAPA COMPLETO DE HOJAS", "ABCDEFGH")
mapa = [
    ("02_METODOLOGIA", "Cómo funciona el archivo, colores, verificación y marco legal — LÉALA PRIMERO"),
    ("08_HISTORICO", "De dónde viene cada cifra: AG2023 → AG2024 → AG2025 conectado en vivo"),
    ("10_MOTOR_EXACTO", "La misma cadena sin redondeo a miles (trazabilidad al peso)"),
    ("11_LIMITE_40", "El techo del 40%/1.340 UVT paso a paso — por qué más certificados no siempre bajan el impuesto"),
    ("13_QA", "Pruebas embebidas de la tarifa + registro de la auditoría Python"),
    ("14_SOPORTES", "La carpeta documental: qué falta y para qué sirve cada papel"),
    ("15_EXOGENA", "Copia de lo que los terceros reportaron a la DIAN"),
    ("16_PARAMETROS", "UVT, topes, tablas y plazos con su norma — el único lugar que cambia cada año"),
    ("17_DECISIONES", "Qué cambió en v4 y con qué evidencia"),
    ("18_GLOSARIO", "Cada término en palabras simples y en precisión técnica"),
]
for i, (h, dsc) in enumerate(mapa):
    r = 29 + i
    put(ws, f"A{r}", f'=HYPERLINK("#\'{h}\'!A1","{h}")', S_fx, fmt="@")
    put(ws, f"B{r}", dsc, S_note); ws.merge_cells(f"B{r}:H{r}")
    ws.row_dimensions[r].height = 15
sec_row(ws, 40, "CONVENCIONES", "ABCDEFGH")
put(ws, "A41", "VERDE = entrada suya", S_in, fmt="@")
put(ws, "B41", "ÁMBAR = ESTIMADO a reemplazar", S_est, fmt="@"); ws.merge_cells("B41:C41")
c = put(ws, "D41", "AZUL = fórmula"); c.font = Font(name=F, size=9, bold=True, color=BLU)
put(ws, "E41", "NARANJA = advertencia", S_warn); ws.merge_cells("E41:F41")
put(ws, "G41", "ROJO = bloqueo", S_alert); ws.merge_cells("G41:H41")
put(ws, "A43", "ADVERTENCIA LEGAL: hoja de trabajo, no dictamen. La coincidencia matemática Excel↔Python NO certifica hechos económicos ni soportes. Responsabilidad del declarante: arts. 571, 596, 647-648 E.T. Valores del formulario aproximados al múltiplo de mil (art. 577).", S_note)
ws.merge_cells("A43:H44")

# ============================================================ orden de pestañas y colores
wb._sheets.sort(key=lambda s: s.title)
tabcol = {SN["INICIO"]: NAVY, SN["MET"]: "44546A", SN["PERFIL"]: GRNB, SN["ING"]: GRNB,
          SN["BEN"]: GRNB, SN["DEP"]: GRNB, SN["PAT"]: GRNB, SN["HIST"]: TEAL,
          SN["F210"]: "2E5395", SN["MX"]: "8496B0", SN["LIM"]: "BF8F00", SN["CTRL"]: "C00000",
          SN["QA"]: "7030A0", SN["SOP"]: "C00000", SN["EXO"]: GRIS, SN["PAR"]: GRIS,
          SN["DEC"]: "7030A0", SN["GLO"]: "375623"}
for n, col in tabcol.items(): wb[n].sheet_properties.tabColor = col
wb.active = 0

out = SALIDA or CFG.get('archivo', 'Declaracion_Renta_F210.xlsx')
wb.save(out)
print("guardado", out)
